"""Build the IIFT Agent/Banca Portal & Back-office tender proposal (.docx).

Usage:
    python docs/proposal/build/build_proposal.py

The document is assembled from section functions. Each section keeps its
wording in plain Python lists/tuples so the text can be edited without
touching the layout helpers in docx_kit.py. All prices come from
pricing_data.py. Diagrams are drawn fresh on every run (diagrams.py), and
screenshots are picked up from docs/proposal/screenshots/ when present.
"""

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import brand                      # noqa: E402
import diagrams                   # noqa: E402
import pricing_data as price      # noqa: E402
from docx_kit import ProposalWriter, add_rich_text   # noqa: E402
from docx.shared import Pt, Cm    # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH        # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT  # noqa: E402

import sections_commercial        # noqa: E402
import sections_delivery          # noqa: E402
import sections_people            # noqa: E402
import sections_scope             # noqa: E402
import sections_solution          # noqa: E402
from numbering import CHAPTERS, sec   # noqa: E402

GANTT_PHASES = [
    ("0  Mobilisation", 1, 2, "build"),
    ("1  Discovery and fit-gap confirmation", 2, 6, "build"),
    ("2  Configuration sprints 1–5", 6, 16, "build"),
    ("3  Integration with IIFT systems", 8, 17, "build"),
    ("4  SIT, performance, VAPT", 16, 19, "build"),
    ("5  UAT, migration rehearsals, training", 19, 23, "build"),
    ("6  Cut-over and go-live", 24, 24, "gate"),
    ("7  Hypercare and closure", 25, 28, "support"),
]
MILESTONES = [
    ("M1", 2, "Project charter & plan approved"),
    ("M2", 6, "Design sign-off"),
    ("M3", 16, "Configuration complete"),
    ("M4", 19, "SIT exit & security clearance"),
    ("M5", 22, "UAT sign-off"),
    ("M6", 24, "Production go-live"),
    ("M7", 28, "Hypercare exit & closure"),
]


# --- Figures ----------------------------------------------------------------------
def render_figures(out_dir: Path) -> dict:
    return {
        "logical": diagrams.logical_architecture(out_dir),
        "on_prem": diagrams.on_prem_deployment(out_dir),
        "cloud": diagrams.cloud_deployment(out_dir),
        "integration": diagrams.integration_flow(out_dir),
        "lifecycle": diagrams.policy_lifecycle(out_dir),
        "governance": diagrams.governance_structure(out_dir),
        "gantt": diagrams.gantt_chart(out_dir, GANTT_PHASES, MILESTONES),
        "cover_band": diagrams.cover_band(out_dir),
        "out_dir": out_dir,
    }


# --- Cover page -----------------------------------------------------------------------
def _wordmark(paragraph, size=26):
    """Text wordmark used when no iorta logo image is supplied."""
    add_rich_text(paragraph, "iorta", size=size, colour=brand.MAGENTA, bold=True)
    add_rich_text(paragraph, " TechNXT", size=size, colour=brand.TEXT_DARK, bold=False)


def cover_page(w: ProposalWriter, figs: dict):
    doc = w.doc
    top = doc.add_paragraph()
    top.paragraph_format.space_after = Pt(30)
    if brand.IORTA_LOGO.exists():
        top.add_run().add_picture(str(brand.IORTA_LOGO), height=Cm(1.6))
    else:
        _wordmark(top)

    w.image(figs["cover_band"], width_cm=17.0)
    w.spacer(18)
    w.para(brand.PROPOSAL_TITLE, size=27, colour=brand.MAGENTA, bold=True, space_after=10)
    w.para(f"{brand.PRODUCT}, configured for IIFT as the {brand.SOLUTION_NAME}", size=14,
           colour=brand.ORANGE, bold=True, space_after=6)
    w.para("Technical and commercial proposal in response to IIFT's Request for Proposal", size=12,
           colour=brand.TEXT_MUTED, space_after=36)

    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    logo_cell, client_cell = table.rows[0].cells
    logo_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    logo_cell.paragraphs[0].add_run().add_picture(str(brand.CLIENT_LOGO), height=Cm(3.6))
    client_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for text, size, colour, bold in (
        ("PREPARED FOR", 9, brand.ORANGE, True),
        (brand.CLIENT, 15, brand.TEXT_DARK, True),
        ("Insurans Islam TAIB Holding (IITH) group", 10, brand.TEXT_MUTED, False),
        ("Bandar Seri Begawan, Brunei Darussalam", 10, brand.TEXT_MUTED, False),
    ):
        paragraph = client_cell.add_paragraph() if client_cell.paragraphs[0].text or text != "PREPARED FOR" else client_cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(3)
        add_rich_text(paragraph, text, size=size, colour=colour, bold=bold)

    spacer_cell, bidder_cell = table.rows[1].cells
    spacer_cell.paragraphs[0].paragraph_format.space_before = Pt(24)
    first = bidder_cell.paragraphs[0]
    first.paragraph_format.space_before = Pt(24)
    for index, (text, size, colour, bold) in enumerate((
        ("SUBMITTED BY", 9, brand.ORANGE, True),
        (brand.BIDDER, 15, brand.TEXT_DARK, True),
        (brand.BIDDER_WEBSITE, 10, brand.TEXT_MUTED, False),
    )):
        paragraph = first if index == 0 else bidder_cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        add_rich_text(paragraph, text, size=size, colour=colour, bold=bold)
    from docx_kit import set_column_widths
    set_column_widths(table, [3.4, 13.6])

    w.spacer(30)
    w.key_value_table([
        ("Submission date", brand.SUBMISSION_DATE),
        ("Proposal reference no.", "[Proposal reference no.]"),
        ("RFP reference", "[IIFT RFP / tender reference no.]"),
        ("Document version", brand.DOCUMENT_VERSION),
        ("Classification", brand.CLASSIFICATION),
    ], widths=(5.0, 12.0))
    w.spacer(16)
    w.para(f"{brand.CLASSIFICATION}. This proposal contains confidential commercial information of "
           f"{brand.BIDDER} and is submitted solely for the evaluation of the above RFP by "
           f"{brand.CLIENT}. It may not be disclosed to third parties without prior written consent.",
           size=8, colour=brand.TEXT_MUTED, italic=True)


# --- Front matter ----------------------------------------------------------------------
def document_control(w: ProposalWriter):
    w.h1("Document Control", numbered=False)
    w.h3("Version history")
    w.table(["Version", "Date", "Author", "Description"],
            [["0.9", "[Date]", brand.BIDDER, "Internal review draft"],
             [brand.DOCUMENT_VERSION, brand.SUBMISSION_DATE, brand.BIDDER, "Submission version"]],
            widths=[2.2, 3.3, 4.0, 7.5])
    w.h3("Review and approval")
    w.table(["Role", "Name", "Signature / date"],
            [["Prepared by (Bid Manager)", "[Name]", ""],
             ["Reviewed by (Solution Architect)", "[Name]", ""],
             ["Approved by (Authorised Signatory)", "[Name, designation]", ""]],
            widths=[6.0, 6.0, 5.0])
    w.h3("Distribution")
    w.table(["Recipient", "Organisation", "Format"],
            [["Tender / Evaluation Committee", brand.CLIENT,
              "Signed PDF and editable DOCX; pricing workbook and Bill of Materials (XLSX)"],
             ["Bid file", brand.BIDDER, "Archive copy"]],
            widths=[6.0, 6.0, 5.0])
    w.h3("Companion documents")
    w.bullets([
        f"{brand.XLSX_OUTPUT.name}: the pricing workbook (options summary, fees in IIFT's format, Options A, B and "
        "C, RFP section 10 breakdown, OPE, rate card, payment schedules and assumptions).",
        f"{brand.BOM_OUTPUT.name}: the Bill of Materials (on-premise infrastructure, cloud services, software "
        "licences and third-party services, with indicative costs and who pays).",
        "Enclosure B: corporate profile and relevant experience (attached separately).",
        "Enclosure C: business registration certificate (attached separately).",
        "Technical document pack in docs/technical: solution architecture, data dictionary, code standards and "
        "quality report, security assessment report, production support handover.",
    ])


def response_map(w: ProposalWriter):
    w.h1("Response Checklist", numbered=False)
    w.para("Where each item (a) to (i) of the invitation is answered, and which items are attached separately.")
    w.table(["Item", "Invitation requirement", "Where addressed", "Status"], [
        ["(a)", "Cover letter", "Cover Letter, after the cover page", "Included"],
        ["(b)", "Corporate profile and relevant experience",
         f"Enclosure B (Section {sec('enclosure_b')})", "To be attached by iorta"],
        ["(c)", "Business registration certificate", f"Enclosure C (Section {sec('enclosure_c')})",
         "To be attached by iorta"],
        ["(d)", "Detailed scope of work and methodology",
         f"Sections {sec('scope')} (scope), {sec('fitment')} (fitment), {sec('methodology')} (methodology)",
         "Included"],
        ["(e)", "Detailed fees and pricing structure, stating WHT, OPE and third-party charges",
         f"Section {sec('commercials')}; pricing workbook; Bill of Materials workbook; Annex C", "Included"],
        ["(f)", "Delivery and implementation timeline with key milestones", f"Section {sec('timeline')}", "Included"],
        ["(g)", "Quotation validity period",
         f"Section {sec('validity')}: {price.QUOTATION_VALIDITY_DAYS} days", "Included"],
        ["(h)", "Other supporting information",
         f"Sections {sec('personas')}, {sec('techdocs')} and {sec('supporting')}; technical document pack",
         "Included"],
        ["(i)", "Terms and conditions", f"Section {sec('terms')}", "Included"],
        ("GROUP", "Additional chapters"),
        ["", "Personas, screens, journeys and navigation",
         f"Section {sec('personas')} (journeys in {sec('journeys')}, screen index in {sec('screens')})", ""],
        ["", "Functional rules and products", f"Sections {sec('functional')} and {sec('products')}", ""],
        ["", "Architecture, security, deployment, infrastructure",
         f"Sections {sec('architecture')} to {sec('infrastructure')}", ""],
        ["", "Technical document pack", f"Section {sec('techdocs')}", ""],
        ["", "Five-year maintenance and support, SLA", f"Section {sec('maintenance')}", ""],
        ["", "Team, assumptions and risks", f"Sections {sec('team')} to {sec('risks')}", ""],
        ["", "Compliance and fitment matrix (every RFP requirement ID)", "Annex A", ""],
        ["", "Bill of Materials summary, software components, network flows", "Annex C", ""],
    ], widths=[1.2, 6.6, 6.0, 3.2], font_size=8, center_cols=(0,))


def table_of_contents(w: ProposalWriter):
    w.h1("Table of Contents", numbered=False)
    w.toc()


ABBREVIATIONS = [
    ("AD / LDAP", "Active Directory / Lightweight Directory Access Protocol, used for single sign-on"),
    ("AMBD", "Autoriti Monetari Brunei Darussalam, the central bank and financial regulator of Brunei Darussalam"),
    ("AMC", "Annual maintenance charge"),
    ("AML / KYC", "Anti-Money Laundering / Know Your Customer screening and due diligence"),
    ("API", "Application Programming Interface"),
    ("ASVS", "OWASP Application Security Verification Standard"),
    ("BRR", "The issuance register / report referenced in RFP Appendix 3 against which daily issuance and FIN postings are tallied"),
    ("BOM / SBOM", "Bill of materials / software bill of materials (open-source components and licences)"),
    ("BRS / FRS", "Business Requirement Specification / Functional Requirement Specification"),
    ("B$ / BND", "Brunei Dollar"),
    ("CI/CD", "Continuous Integration / Continuous Delivery"),
    ("CR", "Change Request"),
    ("DAST", "Dynamic Application Security Testing (scanning the running application)"),
    ("DMZ", "Demilitarised zone: the network segment between the internet and the internal network"),
    ("DR", "Disaster Recovery"),
    ("EOD", "End-of-Day processing"),
    ("FIN", "IIFT financial / accounting system referenced in the RFP"),
    ("FFR", "Functional Flow and Requirements (RFP Appendix 3)"),
    ("HA", "High Availability"),
    ("IC", "Identity Card (Brunei national identity card)"),
    ("IIFT", "Insurans Islam Family Takaful Sendirian Berhad"),
    ("IIGT", "Insurans Islam General Takaful, the sister company that operates the IIGT Agent Portal"),
    ("IITH", "Insurans Islam TAIB Holding, the group holding company"),
    ("ITIL", "Information Technology Infrastructure Library, the reference for service management practice"),
    ("KT", "Knowledge transfer"),
    ("L1 / L2 / L3", "First-, second- and third-line support"),
    ("MFA / OTP / TOTP", "Multi-Factor Authentication / One-Time Password / Time-based One-Time Password"),
    ("mTLS", "Mutual TLS: both ends of a connection present certificates"),
    ("OPE", "Out-of-Pocket Expenses (travel, accommodation, per diem, insurance, transfers, visas)"),
    ("OWASP", "Open Worldwide Application Security Project"),
    ("PDPO", "Personal Data Protection Order 2025 (Brunei Darussalam)"),
    ("PEP", "Politically exposed person (AML watch-list category)"),
    ("PITR", "Point-in-Time Recovery"),
    ("PMO", "Project Management Office"),
    ("QC", "Quality Check"),
    ("RACI", "Responsible, Accountable, Consulted, Informed"),
    ("RBAC", "Role-Based Access Control"),
    ("RPO / RTO", "Recovery Point Objective (maximum data loss) / Recovery Time Objective (maximum downtime)"),
    ("RTM", "Requirements Traceability Matrix"),
    ("SAD / TDD", "Solution Architecture Document / Technical Design Document"),
    ("SAST / SCA", "Static Application Security Testing / Software Composition Analysis"),
    ("SIEM", "Security Information and Event Management (central security log platform)"),
    ("SIT / UAT", "System Integration Testing / User Acceptance Testing"),
    ("SLA", "Service Level Agreement"),
    ("SSO", "Single Sign-On"),
    ("TCO", "Total Cost of Ownership"),
    ("VAPT", "Vulnerability Assessment and Penetration Testing"),
    ("WAF", "Web Application Firewall"),
    ("WAL", "Write-Ahead Log, the PostgreSQL transaction log used for replication and point-in-time recovery"),
    ("WHT", "Withholding Tax"),
]


def abbreviations(w: ProposalWriter):
    w.h1("Abbreviations", numbered=False)
    w.para("Takaful and technical terms that are not abbreviations are explained in the glossary, Annex D.")
    w.table(["Abbreviation", "Meaning"], [list(a) for a in ABBREVIATIONS], widths=[3.6, 13.4],
            bold_first_col=True, font_size=8, padding=30)


# --- (a) Cover letter --------------------------------------------------------------------
COVER_LETTER_BODY = [
    "**Re: Request for Proposal for an Agent/Banca Portal & Back-office Solution, dated [RFP date]**",
    # 1. Submission
    f"We refer to your Request for Proposal dated [RFP date] for an Agent/Banca Portal and Back-office Solution and "
    f"enclose the technical and commercial proposal of {brand.BIDDER}, together with the pricing workbook and the "
    "Bill of Materials. The items (a) to (i) requested in the invitation are mapped in the Response Checklist at "
    "the front of the proposal; the corporate profile (Enclosure B) and the business registration certificate "
    "(Enclosure C) are attached as separate documents.",
    # 2. What is offered
    f"We propose {brand.PRODUCT}, our distribution platform for Takaful and insurance, configured as the "
    f"{brand.SOLUTION_NAME}. The platform already runs the seven products in Appendix 3 of the RFP and the business "
    "rules that go with them, so the project is one of configuration and integration rather than development: "
    "IIFT's rates and templates, the interfaces to the core system, FIN, AML screening, Active Directory, SMS and "
    "e-mail, then testing, migration and training. Go-live is planned for week 24 after kick-off, followed by four "
    "weeks of hypercare, a six-month warranty and five years of maintenance and support. We would welcome the "
    "opportunity to demonstrate the working application to the evaluation committee.",
    # 3. Key commercial facts
    "Two commercial models are offered for the same scope, both at a fixed implementation price. Option A, which we "
    f"recommend, is an on-premise perpetual licence in the IITH data centre: {price.bnd(price.one_time_total('A'))} "
    f"one-time (licence {price.bnd(price.licence_fee())}, services {price.bnd(price.services_fee())}) and an annual "
    f"maintenance charge of {price.bnd(price.amc(1))} in Year 1, rising by 5% a year. Option B is a subscription "
    f"hosted and managed by iorta on cloud: {price.bnd(price.option_b_one_time())} one-time (services "
    f"{price.bnd(price.services_fee())}, cloud set-up {price.bnd(price.CLOUD_SETUP_FEE)}), then "
    f"{price.bnd(price.subscription_monthly(1))} a month for the subscription and "
    f"{price.bnd(price.managed_monthly(1))} a month for managed services in Year 1, with a minimum term of "
    f"{price.SUBSCRIPTION_MINIMUM_MONTHS} months and cloud infrastructure recharged at cost. Option C, an add-on "
    f"under either model, hands over the core platform source code with a knowledge transfer of "
    f"{price.KNOWLEDGE_TRANSFER_WEEKS} weeks for {price.bnd(price.option_c_total())}. All fees are inclusive of withholding tax; "
    f"out-of-pocket expenses (estimated at {price.bnd(price.ope_total())}) and cloud infrastructure are recharged "
    "at cost as disbursements, and third-party charges listed in the Bill of Materials are excluded. The proposal "
    f"is valid for {price.QUOTATION_VALIDITY_DAYS} days from the date of submission.",
    # 4. Closing
    "[Name of contact person], [Designation], is our contact for this tender and can be reached at [e-mail address] "
    "or [telephone number]. We thank you for the invitation and look forward to discussing the proposal with IIFT.",
    "Yours sincerely,",
]


def cover_letter(w: ProposalWriter):
    w.h1("Cover Letter", numbered=False)
    w.para(f"[Day] {brand.SUBMISSION_DATE}", space_after=0)
    w.para("Our ref: [Proposal reference no.]", space_after=0)
    w.para("Your ref: [IIFT RFP / tender reference no.]", space_after=6)
    for line in ("[Name of IIFT procurement contact]", "[Designation]", f"**{brand.CLIENT}**", "[Street address]",
                 "Bandar Seri Begawan [Postcode]", "Brunei Darussalam"):
        w.para(line, space_after=0)
    w.spacer(4)
    w.para("Dear [Salutation and surname],")
    w.paras(COVER_LETTER_BODY)
    w.spacer(2)
    for line in ("______________________________", "[Name of authorised signatory]", "[Designation]",
                 f"For and on behalf of **{brand.BIDDER}**", "[Company stamp]"):
        w.para(line, space_after=0)


# --- Executive summary ---------------------------------------------------------------------
def _available_today():
    """(IDs that run in the platform today, functional and non-functional IDs) from the fitment data."""
    import compliance_matrix as cm
    available = total = 0
    for title, counts, group_total in cm.fitment_summary():
        if title.startswith(("RFP 4", "RFP 5")):
            available += counts["A"]
            total += group_total
    return available, total


def what_iift_receives():
    available, functional_total = _available_today()
    return [
        ("Solution", f"{brand.PRODUCT}, configured as the {brand.SOLUTION_NAME}: one portal for agents and bank "
                     "officers and one back-office for IIFT and IITH staff, on one database and one audit trail."),
        ("Scope", "All 29 components of RFP section 3.1, the seven Appendix 3 products, six interfaces (core system, "
                  "FIN, AML screening, Active Directory, SMS, e-mail), migration, training and documentation."),
        ("Timeline", "Go-live in week 24 from kick-off; hypercare to week 28; six-month warranty; five years of "
                     "maintenance and support."),
        ("Option A (recommended)", f"On-premise perpetual licence in the IITH data centre: "
                                   f"{price.bnd_usd(price.one_time_total('A'))} one-time; AMC "
                                   f"{price.bnd(price.amc(1))} in Year 1 (22% of licence and customisation), +5% a "
                                   "year."),
        ("Option B", f"Subscription on iorta-managed cloud: {price.bnd_usd(price.option_b_one_time())} one-time "
                     f"(including cloud set-up {price.bnd(price.CLOUD_SETUP_FEE)}); subscription "
                     f"{price.bnd(price.subscription_monthly(1))} and managed services "
                     f"{price.bnd(price.managed_monthly(1))} a month in Year 1, +5% a year, minimum "
                     f"{price.SUBSCRIPTION_MINIMUM_MONTHS} months; cloud recharged at cost."),
        ("Option C (add-on)", f"Core platform source code with {price.KNOWLEDGE_TRANSFER_WEEKS} weeks of knowledge "
                              f"transfer after hypercare: {price.bnd(price.option_c_total())}. IIFT-specific source "
                              "code is delivered under every option."),
        ("Five-year cost", f"Option A {price.bnd_usd(price.tco_option_a())}; Option B "
                           f"{price.bnd_usd(price.tco_option_b())}, both including estimated OPE and, for B, cloud "
                           "charges."),
        ("Fee basis", "One-time fees fixed; recurring fees +5% a year, capped; unlimited IIFT users. Inclusive of "
                      "WHT; OPE and cloud recharged at cost as disbursements; third-party charges excluded (Bill of "
                      "Materials)."),
        ("Compliance", f"234 of 234 RFP requirement IDs answered (Annex A). {available} of the {functional_total} "
                       f"functional and non-functional requirements run in the platform today (Section "
                       f"{sec('fitment')})."),
    ]


def _why_low_risk():
    paid_against_deliverables = sum(share for code, _, share in price.SERVICE_MILESTONES if code != "M1")
    return [
        "**The platform is built and can be demonstrated.** IIFT's seven products, the B$ 150,000 referral to IIFT "
        "Sales, the seven-day grace-period block, maker-checker with amount thresholds, watch-list screening, "
        "e-signature, e-Policy and e-Receipt PDFs and the end-of-day FIN interface file run today. Discovery starts "
        "from working screens and takes five weeks, and what IIFT sees in UAT is what it saw in the demonstration. "
        f"What remains (IIFT's rates and templates, the six interfaces, MFA, infrastructure, testing, migration and "
        f"training) is listed in Section {sec('fitment')}.",
        "**Controls are enforced where a screen cannot bypass them.** The API rejects a maker approving their own "
        "request, a database trigger makes the audit table append-only, and IC and passport numbers are encrypted "
        "field by field. An independent penetration test is included before go-live.",
        f"**Commercial exposure is bounded.** Implementation is a fixed price, {paid_against_deliverables:.0%} of the "
        "services fee is paid against signed deliverables, increases are capped at 5% a year and there are no "
        "per-user fees. IIFT can move from subscription to licence during the term, and take the platform source "
        "code once the solution is stable.",
    ]


def executive_summary(w: ProposalWriter):
    w.h1("Executive Summary")
    w.para("IIFT asked for a portal and a back-office that take Banca and agency business out of paper and e-mail "
           f"into one controlled system. We propose {brand.PRODUCT}, configured as the {brand.SOLUTION_NAME}, live in "
           "week 24 at a fixed implementation price with five years of support, and we recommend Option A, a "
           "perpetual licence in the IITH data centre: data stays in Brunei and the five-year cost is the lower of "
           "the two models.")
    w.key_value_table(what_iift_receives(), widths=(3.6, 13.4), caption="What IIFT receives")
    w.h2("Why the delivery risk is low")
    w.paras(_why_low_risk())
    w.para("**Next step.** We invite the evaluation committee to a demonstration of the working application with "
           "IIFT's own products and sample cases, and IITH IT to a session on the two deployment models. The proposal "
           f"is valid for {price.QUOTATION_VALIDITY_DAYS} days (Section {sec('validity')}); the corporate profile and "
           "relevant experience are in Enclosure B.")


# --- Enclosures (b) and (c) -------------------------------------------------------------------
def enclosure_b(w: ProposalWriter):
    w.h1("Enclosure B – Corporate Profile and Relevant Experience")
    w.para("**Attached separately.** iorta TechNXT's latest corporate profile and relevant experience (invitation "
           "item b) are submitted as a separate document with this proposal.")
    w.key_value_table([
        ("Document", "[File name of the corporate profile]"),
        ("Contents", "Company overview, leadership, delivery capability and relevant reference projects"),
        ("Contact", "[Name, designation, e-mail, telephone]"),
    ], caption="Enclosure B: document reference")


def enclosure_c(w: ProposalWriter):
    w.h1("Enclosure C – Business Registration Certificate", new_page=False)
    w.para("**Attached separately.** A certified copy of iorta TechNXT's business registration certificate "
           "(invitation item c) is submitted with this proposal.")
    w.key_value_table([
        ("Document", "[File name of the certificate]"),
        ("Business registration no.", "[Business registration no.]"),
    ], caption="Enclosure C: document reference")


# --- Assembly ----------------------------------------------------------------------------
def chapter_builders(figs: dict) -> dict:
    """Chapter key -> callable, in the order defined by numbering.CHAPTERS."""
    return {
        "executive_summary": lambda w: executive_summary(w),
        "understanding": sections_solution.understanding,
        "enclosure_b": enclosure_b,
        "enclosure_c": enclosure_c,
        "scope": sections_scope.scope_of_work,
        "fitment": sections_scope.fitment_chapter,
        "personas": lambda w: sections_people.people_chapter(w, figs),
        "functional": lambda w: sections_solution.functional_solution(w, figs),
        "products": sections_solution.product_flows,
        "architecture": lambda w: sections_solution.architecture(w, figs),
        "security": sections_solution.security,
        "deployment": sections_solution.deployment,
        "infrastructure": lambda w: sections_solution.infrastructure(w, figs),
        "techdocs": sections_delivery.technical_documents,
        "methodology": lambda w: sections_delivery.methodology(w, figs),
        "timeline": lambda w: sections_delivery.timeline(w, figs, GANTT_PHASES, MILESTONES),
        "commercials": sections_commercial.commercials,
        "maintenance": sections_delivery.maintenance,
        "team": sections_delivery.team,
        "assumptions": sections_delivery.assumptions,
        "risks": sections_delivery.risks,
        "validity": sections_commercial.validity,
        "supporting": sections_commercial.supporting_information,
        "terms": sections_commercial.terms_and_conditions,
    }


def build(output: Path = brand.DOCX_OUTPUT) -> Path:
    price.verify()
    w = ProposalWriter()
    w.set_properties(
        title=brand.PROPOSAL_TITLE,
        subject=f"{brand.PRODUCT} proposal to {brand.CLIENT}",
        keywords=f"IIFT; Takaful; {brand.PRODUCT}; Agent Portal; Bancassurance; Back-office; Proposal",
    )
    w.header_footer()
    with tempfile.TemporaryDirectory() as tmp:
        figs = render_figures(Path(tmp))
        cover_page(w, figs)
        cover_letter(w)
        document_control(w)
        table_of_contents(w)
        response_map(w)
        abbreviations(w)
        builders = chapter_builders(figs)
        for key in CHAPTERS:
            builders[key](w)
            assert w.chapter == int(sec(key)), f"Chapter numbering out of step at '{key}'"
        sections_commercial.annexes(w)
        w.save(output)
    return output


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
