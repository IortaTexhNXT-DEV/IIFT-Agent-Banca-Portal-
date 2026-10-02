"""Shared styling, layout and verification helpers for the client workbooks
(pricing workbook and bill of materials).

SheetWriter keeps a row cursor so sheets can be laid out top to bottom.
FormulaEvaluator recomputes the formulas these scripts write (SUM, ROUND,
arithmetic and cross-sheet references) so the build can check every total
against pricing_data.py without a spreadsheet engine. recalc_errors() runs
the workbook through LibreOffice, when available, and reports any cell that
evaluates to an error.
"""

import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

import brand

HEADER_FILL = PatternFill("solid", fgColor=brand.MAGENTA)
TOTAL_FILL = PatternFill("solid", fgColor=brand.GROUP_ROW)
GROUP_FILL = PatternFill("solid", fgColor="FDF6EF")
ZEBRA_FILL = PatternFill("solid", fgColor=brand.ZEBRA)
INPUT_FILL = PatternFill("solid", fgColor="FFF9DB")
HEADER_FONT = Font(name=brand.BODY_FONT, bold=True, color=brand.WHITE, size=10)
TITLE_FONT = Font(name=brand.BODY_FONT, bold=True, color=brand.MAGENTA, size=15)
SECTION_FONT = Font(name=brand.BODY_FONT, bold=True, color=brand.ORANGE, size=11)
BODY_FONT = Font(name=brand.BODY_FONT, color=brand.TEXT_DARK, size=10)
BOLD_FONT = Font(name=brand.BODY_FONT, color=brand.TEXT_DARK, size=10, bold=True)
NOTE_FONT = Font(name=brand.BODY_FONT, color=brand.TEXT_MUTED, size=9, italic=True)
THIN = Side(style="thin", color=brand.BORDER_GREY)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0;(#,##0);"–"'
MONEY_2DP = '#,##0.00;(#,##0.00);"–"'
PERCENT = "0%"
WRAP = Alignment(wrap_text=True, vertical="top")
ERROR_PREFIXES = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "Err:")


class SheetWriter:
    """Writes a sheet top to bottom with consistent styling."""

    def __init__(self, ws, columns):
        self.ws = ws
        self.columns = columns
        self.row = 1

    # -- blocks -------------------------------------------------------------------
    def title(self, text, subtitle=None, row=None):
        self.row = row or self.row
        self.ws.cell(row=self.row, column=1, value=text).font = TITLE_FONT
        self.ws.row_dimensions[self.row].height = 22
        self.row += 1
        subtitle = subtitle or f"{brand.BIDDER} – proposal to {brand.CLIENT} – {brand.SUBMISSION_DATE}"
        self.ws.cell(row=self.row, column=1, value=subtitle).font = NOTE_FONT
        self.row += 2

    def section(self, text):
        self.ws.cell(row=self.row, column=1, value=text).font = SECTION_FONT
        self.ws.row_dimensions[self.row].height = 18
        self.row += 1

    def _merge(self, row, start, end, fill=None):
        """Merge columns start..end of a row into one bordered cell (for notes)."""
        for col in range(start, end + 1):
            cell = self.ws.cell(row=row, column=col)
            cell.border = BORDER
            if fill:
                cell.fill = fill
        self.ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)

    def header(self, headers, height=30, merge_from=None):
        for col, text in enumerate(headers, start=1):
            cell = self.ws.cell(row=self.row, column=col, value=text)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        if merge_from:
            start, end = merge_from if isinstance(merge_from, tuple) else (merge_from, self.columns)
            self._merge(self.row, start, end, HEADER_FILL)
        self.ws.row_dimensions[self.row].height = height
        self.row += 1
        return self.row - 1

    def line(self, values, money=(), money2=(), pct=(), center=(), bold=False, total=False, zebra=False,
             group=False, inputs=(), height=None, merge_from=None):
        """Write one table row; columns are 1-based. Returns the row number."""
        r = self.row
        fill = TOTAL_FILL if total else (GROUP_FILL if group else (ZEBRA_FILL if zebra else None))
        for col, value in enumerate(values, start=1):
            cell = self.ws.cell(row=r, column=col, value=value)
            cell.font = BOLD_FONT if (bold or total or group) else BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP
            if col in money or col in money2:
                cell.number_format = MONEY_2DP if col in money2 else MONEY
                cell.alignment = Alignment(horizontal="right", vertical="top")
            elif col in pct:
                cell.number_format = PERCENT
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col in center:
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            if fill:
                cell.fill = fill
            if col in inputs:
                cell.fill = INPUT_FILL
        if merge_from:
            start, end = merge_from if isinstance(merge_from, tuple) else (merge_from, self.columns)
            self._merge(r, start, end, fill)
        if height:
            self.ws.row_dimensions[r].height = height
        self.row += 1
        return r

    def group(self, text):
        """Full-width group row inside a table."""
        r = self.line([text] + [None] * (self.columns - 1), group=True)
        return r

    def note(self, text, columns=None, height=None):
        columns = columns or self.columns
        cell = self.ws.cell(row=self.row, column=1, value=text)
        cell.font = NOTE_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=columns)
        chars_per_line = max(40, int(sum(self.ws.column_dimensions[get_column_letter(c)].width or 10
                                         for c in range(1, columns + 1)) * 1.1))
        self.ws.row_dimensions[self.row].height = height or 13 * (1 + len(text) // chars_per_line)
        self.row += 1

    def bullets(self, items, columns=None):
        for item in items:
            self.note(f"•  {item}", columns)
            self.ws.cell(row=self.row - 1, column=1).font = BODY_FONT

    def label_value(self, label, value, number_format=None, input_cell=False, value_col=3):
        self.ws.cell(row=self.row, column=1, value=label).font = BOLD_FONT
        cell = self.ws.cell(row=self.row, column=value_col, value=value)
        cell.font = BODY_FONT
        cell.alignment = Alignment(horizontal="left", vertical="top")
        if number_format:
            cell.number_format = number_format
        if input_cell:
            cell.fill = INPUT_FILL
            cell.border = BORDER
        self.row += 1
        return cell

    def skip(self, rows=1):
        self.row += rows


def set_widths(ws, values):
    for index, width in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def add_logo(ws, anchor="A1", height_px=52):
    """Place the iorta TechNXT logo; returns True if the logo file exists."""
    if not brand.IORTA_LOGO.exists():
        return False
    image = XLImage(str(brand.IORTA_LOGO))
    ratio = height_px / image.height
    image.height = height_px
    image.width = int(image.width * ratio)
    ws.add_image(image, anchor)
    return True


def finish(wb, title, subject, first_sheet):
    """Common sheet settings: tab colours, gridlines, A4 landscape fit to width, footer."""
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = brand.MAGENTA if ws is first_sheet else brand.ORANGE
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.6
        ws.oddHeader.right.text = f"{brand.BIDDER} | {brand.SOLUTION_NAME}"
        ws.oddHeader.right.size = 8
        ws.oddFooter.left.text = brand.CLASSIFICATION
        ws.oddFooter.left.size = 8
        ws.oddFooter.right.text = "&A | Page &P of &N"
        ws.oddFooter.right.size = 8
    props = wb.properties
    props.creator = brand.BIDDER
    props.lastModifiedBy = brand.BIDDER
    props.title = title
    props.subject = subject
    props.company = brand.BIDDER
    wb.active = 0


# --- Verification --------------------------------------------------------------------
class FormulaEvaluator:
    """Evaluates the subset of Excel formulas written by the build scripts."""

    REF = r"(?:'([^']+)'!|([A-Za-z_][\w ]*)!)?\$?([A-Z]{1,3})\$?(\d+)"

    def __init__(self, path: Path):
        self.wb = load_workbook(str(path))
        self.cache = {}

    def value(self, sheet, ref):
        key = (sheet, ref.replace("$", ""))
        if key not in self.cache:
            raw = self.wb[sheet][key[1]].value
            if isinstance(raw, str) and raw.startswith("="):
                self.cache[key] = self._evaluate(sheet, raw)
            elif isinstance(raw, (int, float)):
                self.cache[key] = raw
            else:
                self.cache[key] = 0
        return self.cache[key]

    def _range(self, sheet, start, end):
        ws = self.wb[sheet]
        return [self.value(sheet, c.coordinate) for row in ws[start:end] for c in row]

    def _evaluate(self, sheet, formula):
        expr = formula[1:]

        def sheet_of(m):
            return m.group(1) or m.group(2) or sheet

        expr = re.sub(r"SUM\(" + self.REF + r":\$?([A-Z]{1,3})\$?(\d+)\)",
                      lambda m: repr(sum(self._range(sheet_of(m), f"{m.group(3)}{m.group(4)}",
                                                     f"{m.group(5)}{m.group(6)}"))), expr)
        expr = re.sub(self.REF, lambda m: repr(self.value(sheet_of(m), f"{m.group(3)}{m.group(4)}")), expr)
        expr = expr.replace("ROUND", "_round").replace("^", "**")
        return eval(expr, {"_round": lambda x, d: round(x + (1e-9 if x >= 0 else -1e-9), int(d))})


def recalc_errors(path: Path) -> list:
    """Recalculate the workbook with LibreOffice and list cells that evaluate to an error.

    Returns [] when LibreOffice is not installed (the Python evaluator still checks totals)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return []
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(path)],
                       check=True, capture_output=True, timeout=180)
        recalculated = load_workbook(str(Path(tmp) / path.name), data_only=True)
        errors = []
        for ws in recalculated.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith(ERROR_PREFIXES):
                        errors.append(f"{ws.title}!{cell.coordinate}: {cell.value}")
        return errors


def recalculated_values(path: Path):
    """Workbook with values computed by LibreOffice (None if LibreOffice is unavailable)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    tmp = tempfile.mkdtemp()
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(path)],
                   check=True, capture_output=True, timeout=180)
    return load_workbook(str(Path(tmp) / path.name), data_only=True)
