"""Build the Bill of Materials workbook (.xlsx) for the IIFT proposal.

Usage:
    python3 docs/proposal/build/build_bom.py

Sheets: Summary, On-premise infrastructure (Option A, IIFT procures), Cloud
services (Option B, at actuals), Software & licences, Third-party services,
Assumptions. All data comes from pricing_data.py; the on-premise sizing is
checked against the Solution Architecture source in docs/technical so the two
documents cannot drift apart. Totals are Excel formulas and are verified
against pricing_data.py after the file is written.
"""

import json
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook       # noqa: E402

import brand                        # noqa: E402
import pricing_data as price        # noqa: E402
from xlsx_kit import (FormulaEvaluator, SheetWriter, add_logo, finish, recalc_errors,  # noqa: E402
                      recalculated_values, set_widths)

REPO = brand.PROPOSAL_DIR.parent.parent
SAD_SOURCE = REPO / "docs" / "technical" / "build" / "sad_part2.py"
LICENCE_EVIDENCE = REPO / "docs" / "technical" / "evidence" / "runtime-dependency-licences.json"
API_PACKAGE = REPO / "apps" / "api" / "package.json"
WEB_PACKAGE = REPO / "apps" / "web" / "package.json"
YEARS = price.CONTRACT_YEARS
REFS = {}


def ref(sheet, col, row):
    return f"'{sheet}'!${col}${row}"


def check_sizing_against_architecture():
    """Every on-premise row must appear verbatim in the Solution Architecture sizing table."""
    source = SAD_SOURCE.read_text(encoding="utf-8")
    missing = []
    for _, server, qty, _, cpu, ram, disk, software, *_ in price.ON_PREM_BOM:
        if qty == "Not required":
            continue
        row = f'["{server}", "{qty}", "{cpu}", "{ram}", "{disk}", "{software}"]'
        if row not in source:
            missing.append(row)
    if missing:
        raise AssertionError("BOM sizing differs from the Solution Architecture:\n  " + "\n  ".join(missing))


def ram_gb(text):
    return int(text.split()[0]) if text and text[0].isdigit() else None


# =============================================================================
def infrastructure_sheet(wb):
    name = "On-premise infrastructure"
    ws = wb.create_sheet(name)
    set_widths(ws, [16, 28, 15, 8, 10, 16, 40, 10, 12, 13, 13, 13, 13, 40])
    s = SheetWriter(ws, 14)
    s.title("On-premise infrastructure for Option A – procured by IIFT to iorta's sizing")
    s.note("Sizing as in the Solution Architecture (Infrastructure Sizing). Costs are indicative one-time amounts if "
           "capacity is provisioned new, for IIFT's budgeting; they are lower where existing IITH virtualisation, "
           "storage, backup and security capacity is reused. They are not part of iorta's fees.", columns=14)
    s.skip()
    s.header(["Environment", "Server / item", "Qty", "vCPU", "RAM (GB)", "Storage", "Software", "OS licence",
              "Total vCPU", "Unit cost low (B$)", "Unit cost high (B$)", "Total low (B$)", "Total high (B$)",
              "Notes"], height=32)
    first = s.row
    os_rows = []
    for i, (grp, server, qty, qty_n, cpu, ram, disk, software, os_vm, (low, high), note) in enumerate(price.ON_PREM_BOM):
        r = s.row
        cpu_n = int(cpu) if cpu.isdigit() else None
        s.line([grp, server, qty_n, cpu_n, ram_gb(ram), disk, software, "Yes" if os_vm else "No",
                f"=C{r}*D{r}" if cpu_n else None, low, high, f"=C{r}*J{r}", f"=C{r}*K{r}",
                (f"Quantity: {qty}. " if not qty.isdigit() and qty != "–" else "") + note],
               money=(10, 11, 12, 13), center=(3, 4, 5, 8, 9), zebra=i % 2 == 1)
        if os_vm:
            os_rows.append(r)
    last = s.row - 1
    r = s.line(["Total", "", f"=SUM(C{first}:C{last})", "", f"=SUM(E{first}:E{last})", "", "", "",
                f"=SUM(I{first}:I{last})", "", "", f"=SUM(L{first}:L{last})", f"=SUM(M{first}:M{last})",
                "Indicative one-time cost, IIFT's own procurement"], money=(12, 13), center=(3, 5, 9), total=True)
    REFS["infra_low"], REFS["infra_high"] = ref(name, "L", r), ref(name, "M", r)
    s.skip()
    s.section("Operating system")
    vms = s.line(["VMs needing an OS licence", "", "=" + "+".join(f"C{x}" for x in os_rows), "", "", "", "", "", "",
                  "", "", "", "", "Application, database, monitoring, SIT, UAT and DR VMs"], center=(3,))
    s.line(["Ubuntu Server 24.04 LTS", "Recommended", None, "", "", "", "No licence fee", "", "", 0, 0,
            f"=C{vms}*J{s.row}", f"=C{vms}*K{s.row}", "Optional Ubuntu Pro support not required"],
           money=(10, 11, 12, 13))
    r = s.line(["RHEL 9 subscriptions", "Optional, per VM per year", None, "", "", "", "Only if IIFT standardises on "
                "RHEL", "", "", price.RHEL_PER_VM_YEAR[0], price.RHEL_PER_VM_YEAR[1], f"=C{vms}*J{s.row}",
                f"=C{vms}*K{s.row}", "Annual; IIFT procures"], money=(10, 11, 12, 13))
    REFS["rhel_low"], REFS["rhel_high"] = ref(name, "L", r), ref(name, "M", r)
    s.skip()
    s.section("Reused IITH services (no new procurement expected)")
    s.bullets([
        "Virtualisation platform, network, firewall, DNS and the DMZ.",
        "WAF / load balancer (F5 or equivalent): reuse recommended; otherwise an HAProxy/nginx pair on two small VMs.",
        "Enterprise backup platform and off-site copies; pgBackRest (open source) handles PostgreSQL backups.",
        "SIEM and infrastructure monitoring; iorta configures log and metric forwarding.",
        "DEV environment: hosted by iorta, so IIFT provides none.",
    ], columns=14)
    ws.freeze_panes = "C6"


def cloud_sheet(wb):
    name = "Cloud services (Option B)"
    ws = wb.create_sheet(name)
    set_widths(ws, [6, 46, 32, 54, 14, 14, 44])
    s = SheetWriter(ws, 7)
    s.title(f"Cloud services for Option B – recharged at cost ({price.CLOUD_REGION})")
    s.note("Estimate for budgeting. iorta recharges the provider's charges monthly in arrears at cost as a "
           "disbursement, with the provider invoice attached and no mark-up, outside the fees to which WHT applies. "
           "Alternatively IIFT holds the cloud account in its own name and pays the provider directly. The one-off "
           f"cloud set-up (B$ {price.CLOUD_SETUP_FEE:,}) and iorta's managed services are fees in the pricing "
           "workbook.", columns=7)
    s.skip()
    s.header(["No", "Service", "Purpose", "Reference sizing", "B$ per month", "B$ per year", "Billing basis"])
    first = s.row
    for i, ((service, purpose, monthly), sizing) in enumerate(zip(price.CLOUD_MONTHLY_ITEMS, price.CLOUD_SIZING),
                                                               start=1):
        r = s.row
        s.line([i, service, purpose, sizing, monthly, f"=E{r}*12", price.CLOUD_BILLING_BASIS], money=(5, 6),
               center=(1,), zebra=i % 2 == 0)
    last = s.row - 1
    r = s.line(["", "Total, production run", "", "", f"=SUM(E{first}:E{last})", f"=SUM(F{first}:F{last})", ""],
               money=(5, 6), total=True)
    REFS["cloud_month"], REFS["cloud_year"] = ref(name, "E", r), ref(name, "F", r)
    p = s.line(["", "Project environments (DEV, SIT, UAT) before go-live", "Implementation period",
                f"Scaled down, {price.CLOUD_IMPLEMENTATION_MONTHS} months", price.CLOUD_IMPLEMENTATION_MONTHLY,
                f"=E{s.row}*{price.CLOUD_IMPLEMENTATION_MONTHS}", "One-off total for the project period"],
               money=(5, 6))
    REFS["cloud_project"] = ref(name, "F", p)
    r = s.line(["", "Estimate, project plus five years", "", "", None, f"=F{p}+{YEARS}*F{r}", ""], money=(6,),
               total=True)
    REFS["cloud_total"] = ref(name, "F", r)
    s.skip()
    s.bullets([
        "Region: AWS Asia Pacific (Malaysia) ap-southeast-5 or Azure Malaysia West; DR copies in a second region "
        "subject to IIFT approval. A Brunei-hosted alternative can be priced on request.",
        "Providers bill in USD; charges are converted to B$ at the bank selling rate on the provider's invoice date.",
        "Budget alerts at 80% and 100% of the estimate; monthly cost report; quarterly right-sizing review.",
    ], columns=7)
    ws.freeze_panes = "C6"


def software_sheet(wb):
    name = "Software & licences"
    ws = wb.create_sheet(name)
    set_widths(ws, [44, 22, 44, 30, 14, 34])
    s = SheetWriter(ws, 6)
    s.title("Software components and licences")
    s.note("Every runtime component of the solution. All are open source with no licence fee; the operating system "
           "is the only item with a paid choice (RHEL).", columns=6)
    s.skip()
    s.section("1. Runtime components")
    s.header(["Component", "Version", "Role", "Licence", "Licence fee (B$)", "Treatment"])
    first = s.row
    for i, (component, version, role, licence) in enumerate(price.SOFTWARE_COMPONENTS):
        treatment = ("Option A: IIFT procures (Ubuntu free; RHEL optional). Option B: in cloud actuals"
                     if component.startswith("Linux") else price.INCLUDED)
        s.line([component, version, role, licence, 0, treatment], money=(5,), zebra=i % 2 == 1)
    s.line(["Total licence fees", "", "", "", f"=SUM(E{first}:E{s.row - 1})", ""], money=(5,), total=True)
    s.skip()

    data = json.loads(LICENCE_EVIDENCE.read_text(encoding="utf-8"))
    counts = Counter(data["byLicense"])
    s.section(f"2. Application packages by licence ({data['total']} runtime npm packages, from the CI licence scan)")
    s.header(["Licence", "Packages", "Type"], merge_from=(3, 6))
    first = s.row
    permissive = {"MIT", "ISC", "Apache-2.0", "BSD-3-Clause", "BSD-2-Clause", "0BSD", "MIT-0", "Unlicense",
                  "MIT/X11", "MIT and ISC", "MIT AND ISC", "(MIT AND Zlib)", "Python-2.0"}
    for i, (licence, count) in enumerate(counts.most_common()):
        if licence in permissive:
            kind = "Permissive"
        elif licence == "UNKNOWN":
            kind = "Not declared in package metadata; reviewed in the SBOM licence check before go-live"
        elif licence.startswith("(MIT OR"):
            kind = "Dual licence; used under MIT"
        else:
            kind = "Weak copyleft; used unmodified as a library"
        s.line([licence, count, kind], center=(2,), zebra=i % 2 == 1, merge_from=(3, 6))
    s.line(["Total", f"=SUM(B{first}:B{s.row - 1})", ""], center=(2,), total=True, merge_from=(3, 6))
    REFS["pkg_total"] = ref(name, "B", s.row - 1)
    names = {p["license"]: [] for p in data["packages"]}
    for p in data["packages"]:
        names[p["license"]].append(f"{p['name']} {p['version']}")
    flagged = [f"{lic}: {', '.join(sorted(set(names[lic])))}" for lic in counts
               if lic not in permissive]
    s.skip()
    s.section("3. Notes")
    api = json.loads(API_PACKAGE.read_text(encoding="utf-8"))
    web = json.loads(WEB_PACKAGE.read_text(encoding="utf-8"))
    s.bullets([
        f"Direct runtime dependencies: {len(api['dependencies'])} in the API (apps/api/package.json) and "
        f"{len(web['dependencies'])} in the web application (apps/web/package.json).",
        "Packages that are not plainly permissive: " + "; ".join(flagged) + ".",
        "ClamAV (GPL-2.0) and Grafana / Loki (AGPL-3.0) run as separate, unmodified programs; IITH monitoring tools "
        "can replace Grafana and Loki.",
        "The CI pipeline produces a CycloneDX SBOM (npm sbom, production dependencies) on every build; it is "
        "delivered with each release.",
    ], columns=6)
    ws.freeze_panes = "A6"


def third_party_sheet(wb):
    name = "Third-party services"
    ws = wb.create_sheet(name)
    set_widths(ws, [40, 20, 11, 22, 12, 12, 22, 12, 12, 70])
    s = SheetWriter(ws, 10)
    s.title("Third-party services, paid and optional items – indicative B$ per year")
    s.note("Treatment: Included in iorta fee / Pass-through at cost / IIFT procures / Not required. Amounts are "
           "indicative and confirmed by each provider; they are not part of iorta's fees unless marked included.",
           columns=10)
    s.skip()
    s.header(["Item", "Basis", "Required", "Option A treatment", "A: low (B$)", "A: high (B$)", "Option B treatment",
              "B: low (B$)", "B: high (B$)", "Notes"], height=32)
    items = sorted(price.THIRD_PARTY_ITEMS, key=lambda item: not item[5])

    def costs(item, option):
        treatment = item[2] if option == "A" else item[3]
        if treatment not in (price.IIFT_PROCURES, price.PASS_THROUGH):
            return 0, 0
        if option == "B" and item[0].startswith("SSL"):
            return 0, 0
        return price.third_party_cost(item)

    rows = {True: [], False: []}
    for i, item in enumerate(items):
        a, b = costs(item, "A"), costs(item, "B")
        r = s.line([item[0], item[1], "Yes" if item[5] else "Optional", item[2], a[0], a[1], item[3], b[0], b[1],
                    item[6]], money=(5, 6, 8, 9), center=(3,), zebra=i % 2 == 1)
        rows[item[5]].append(r)
    for required, label in ((True, "Total required items, per year"), (False, "Total optional items, per year")):
        f, l = rows[required][0], rows[required][-1]
        r = s.line([label, "", "", "", f"=SUM(E{f}:E{l})", f"=SUM(F{f}:F{l})", "", f"=SUM(H{f}:H{l})",
                    f"=SUM(I{f}:I{l})", ""], money=(5, 6, 8, 9), total=True)
        key = "req" if required else "opt"
        for col, k in (("E", "a_low"), ("F", "a_high"), ("H", "b_low"), ("I", "b_high")):
            REFS[f"tp_{key}_{k}"] = ref(name, col, r)
    ws.freeze_panes = "B6"


def summary_sheet(ws):
    ws.title = "Summary"
    set_widths(ws, [54, 26, 14, 14, 14, 14, 15, 15])
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 16
    add_logo(ws, "A1", height_px=56)
    s = SheetWriter(ws, 8)
    s.title(f"{brand.PROPOSAL_TITLE}: Bill of Materials", row=4)
    s.label_value("Client", brand.CLIENT, value_col=2)
    s.label_value("Bidder", brand.BIDDER, value_col=2)
    s.label_value("Solution", f"{brand.PRODUCT}, configured as the {brand.SOLUTION_NAME}", value_col=2)
    s.label_value("Currency", "Brunei Dollar (B$), indicative amounts", value_col=2)
    s.skip()
    headers = ["Category", "Treatment", "One-time low", "One-time high", "Per year low", "Per year high",
               "Five years low", "Five years high"]

    def row(label, treatment, one_low, one_high, year_low, year_high, total=False):
        r = s.row
        return s.line([label, treatment, one_low, one_high, year_low, year_high, f"=C{r}+{YEARS}*E{r}",
                       f"=D{r}+{YEARS}*F{r}"], money=(3, 4, 5, 6, 7, 8), total=total)

    def totals(label, rows):
        r = s.row
        return s.line([label, ""] + [f"=" + "+".join(f"{c}{x}" for x in rows) for c in "CDEFGH"],
                      money=(3, 4, 5, 6, 7, 8), total=True)

    s.section("Option A – on-premise (B$)")
    s.header(headers)
    a = [row("On-premise infrastructure (servers, storage, DR)", price.IIFT_PROCURES, f"={REFS['infra_low']}",
             f"={REFS['infra_high']}", 0, 0),
         row("Operating system (Ubuntu Server LTS)", price.IIFT_PROCURES, 0, 0, 0, 0),
         row("Runtime software licences (open source)", price.INCLUDED, 0, 0, 0, 0),
         row("Third-party services – required (SMS, AML data, certificates, e-mail)", "IIFT procures / at cost", 0,
             0, f"={REFS['tp_req_a_low']}", f"={REFS['tp_req_a_high']}")]
    ra = totals("Total Option A, required items", a)
    opt_a = row("Optional items (PostgreSQL support, RHEL, escrow)", "Only if IIFT elects", 0, 0,
                f"={REFS['tp_opt_a_low']}", f"={REFS['tp_opt_a_high']}")
    REFS["sum_a"] = ra
    s.skip()
    s.section("Option B – iorta-hosted cloud (B$)")
    s.header(headers)
    b = [row("Cloud services at cost (estimate)", "Disbursement at cost or paid by IIFT directly", f"={REFS['cloud_project']}",
             f"={REFS['cloud_project']}", f"={REFS['cloud_year']}", f"={REFS['cloud_year']}"),
         row("Runtime software licences (open source)", price.INCLUDED, 0, 0, 0, 0),
         row("Third-party services – required (SMS, AML data, e-mail)", "IIFT procures", 0, 0,
             f"={REFS['tp_req_b_low']}", f"={REFS['tp_req_b_high']}")]
    rb = totals("Total Option B, required items", b)
    row("Optional items (escrow)", "Only if IIFT elects", 0, 0, f"={REFS['tp_opt_b_low']}",
        f"={REFS['tp_opt_b_high']}")
    REFS["sum_b"] = rb
    s.skip()
    s.section("Notes")
    s.bullets([
        "Infrastructure and third-party items are paid by IIFT directly or passed through by iorta at cost, without "
        "mark-up. They are not included in iorta's fees in the pricing workbook.",
        "Independent VAPT and re-test before go-live, the monitoring stack and the CI/CD tooling are included in "
        "iorta's fees.",
        f"Under Option B, the one-off cloud set-up fee (B$ {price.CLOUD_SETUP_FEE:,}) and iorta's managed services "
        "are iorta fees in the pricing workbook, not in this BOM. Cloud charges are recharged at cost as "
        "disbursements with provider invoices attached, or paid by IIFT directly if it holds the account.",
        f"Sizing basis: 26 named users, about 628 policies a year; designed for 100 concurrent and 500 named users. "
        f"Ranges are indicative for budgeting.",
    ], columns=8)
    ws.freeze_panes = "A10"
    return opt_a


def assumptions_sheet(wb):
    ws = wb.create_sheet("Assumptions")
    set_widths(ws, [6, 150])
    s = SheetWriter(ws, 2)
    s.title("Assumptions and notes")
    s.header(["#", "Assumption"], height=20)
    items = [
        "Option A sizing is taken from the Solution Architecture (Infrastructure Sizing) and is checked automatically "
        "against it when the BOM is built.",
        "Indicative infrastructure costs assume new virtual capacity and storage on the IITH platform; reuse of "
        "existing capacity reduces them. Physical hardware, if required, is quoted by IIFT's suppliers.",
        "Ubuntu Server 24.04 LTS is the baseline operating system (no licence fee); RHEL 9 is supported if IIFT "
        "prefers it, at IIFT's cost.",
        "SMS volume of about 6,000 messages a year (OTP and alerts) at B$0.05–0.10 per message.",
        "AML screening reuses IIFT's existing service or lists where possible; a new subscription is priced by the "
        "provider.",
        "Cloud estimates use public list prices for the Malaysia region converted at USD 1 = B$ 1.30; actual charges "
        "vary with usage and exchange rates.",
        "Escrow is an alternative to Option C (source code handover); agent fees are paid by IIFT.",
        "All amounts exclude iorta's fees, which are in the pricing workbook.",
    ]
    for i, text in enumerate(items, start=1):
        s.line([i, text], center=(1,), zebra=i % 2 == 0)


def build(output: Path = brand.BOM_OUTPUT) -> Path:
    price.verify()
    check_sizing_against_architecture()
    REFS.clear()
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    infrastructure_sheet(wb)
    cloud_sheet(wb)
    software_sheet(wb)
    third_party_sheet(wb)
    assumptions_sheet(wb)
    summary_sheet(summary)
    finish(wb, f"{brand.PROPOSAL_TITLE}: Bill of Materials", f"Bill of materials for {brand.CLIENT}", summary)
    summary.page_setup.fitToHeight = 1
    wb.save(str(output))
    return output


def verify(path: Path) -> None:
    ev = FormulaEvaluator(path)
    infra_low, infra_high = price.bom_onprem_totals()
    tp_a, tp_b = price.third_party_annual("A"), price.third_party_annual("B")
    opt_a, opt_b = price.third_party_annual("A", True), price.third_party_annual("B", True)
    rhel = price.rhel_annual()
    data = json.loads(LICENCE_EVIDENCE.read_text(encoding="utf-8"))

    def at(r):
        sheet, cell = r.rsplit("!", 1)
        return sheet.strip("'"), cell.replace("$", "")

    checks = {
        "infrastructure low": (REFS["infra_low"], infra_low),
        "infrastructure high": (REFS["infra_high"], infra_high),
        "RHEL low": (REFS["rhel_low"], rhel[0]),
        "RHEL high": (REFS["rhel_high"], rhel[1]),
        "cloud per month": (REFS["cloud_month"], price.cloud_monthly()),
        "cloud five-year estimate": (REFS["cloud_total"], price.cloud_total()),
        "third-party A low": (REFS["tp_req_a_low"], tp_a[0]),
        "third-party A high": (REFS["tp_req_a_high"], tp_a[1]),
        "third-party B low": (REFS["tp_req_b_low"], tp_b[0]),
        "third-party B high": (REFS["tp_req_b_high"], tp_b[1]),
        "optional A high": (REFS["tp_opt_a_high"], opt_a[1]),
        "optional B high": (REFS["tp_opt_b_high"], opt_b[1]),
        "packages": (REFS["pkg_total"], data["total"]),
        "Summary A five-year high": (f"'Summary'!$H${REFS['sum_a']}", infra_high + YEARS * tp_a[1]),
        "Summary B five-year high": (f"'Summary'!$H${REFS['sum_b']}",
                                     price.cloud_total() + YEARS * tp_b[1]),
    }
    failures = []
    values = recalculated_values(path)
    for name, (cell_ref, expected) in checks.items():
        sheet, cell = at(cell_ref)
        actual = ev.value(sheet, cell)
        if actual != expected:
            failures.append(f"{name}: {actual} != {expected}")
        if values is not None and values[sheet][cell].value != expected:
            failures.append(f"LibreOffice {name}: {values[sheet][cell].value} != {expected}")
    errors = recalc_errors(path)
    if errors:
        failures.append(f"LibreOffice errors: {errors[:10]}")
    print(f"  {len(checks)} BOM totals checked (Python evaluator"
          f"{' and LibreOffice' if values is not None else ''}); sizing matches the Solution Architecture")
    if failures:
        raise AssertionError("BOM check failed:\n  " + "\n  ".join(failures))


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
    verify(path)
    print("All BOM totals verified.")
