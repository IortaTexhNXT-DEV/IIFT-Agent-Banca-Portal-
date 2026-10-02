"""Small helper layer over python-docx for a branded tender document.

ProposalWriter wraps a python-docx Document and offers one method per
building block (headings, paragraphs, bullet lists, tables, figures,
call-outs). Text passed to these helpers supports two inline conventions:

    **bold text**          rendered bold
    [Placeholder text]     rendered with yellow highlight, to be filled in

Headings, figures and tables are numbered automatically.
"""

import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

import brand

INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\])")
TWIPS_PER_CM = 567


# --- Low-level XML helpers ---------------------------------------------------
def _set_font(run_or_style_font, element, name=brand.BODY_FONT):
    """Set a font for all scripts (python-docx only sets ascii/hAnsi)."""
    run_or_style_font.name = name
    rpr = element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(theme_attr), None)


def shade_cell(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    for existing in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def _set_cell_borders(cell, **edges):
    """edges: top/bottom/left/right = (size_eighths_pt, hex) or None for 'nil'."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, spec in edges.items():
        el = OxmlElement(f"w:{edge}")
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            size, colour = spec
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), colour)
        borders.append(el)


def _table_borders(table, colour=brand.BORDER_GREY, size=4, inside=True):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    edges = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for edge in edges:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), colour)
        borders.append(el)
    tbl_pr.append(borders)


def _no_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _cell_margins(table, top=50, bottom=50, left=90, right=90):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def set_column_widths(table, widths_cm):
    """Fix column widths so both Word and LibreOffice honour them."""
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths_cm) * TWIPS_PER_CM)))
    tbl_w.set(qn("w:type"), "dxa")
    grid_cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
    for grid_col, width in zip(grid_cols, widths_cm):
        grid_col.set(qn("w:w"), str(int(width * TWIPS_PER_CM)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def _row_flags(row, header=False):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    if header:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)


def _paragraph_border(paragraph, edge="bottom", colour=brand.ORANGE, size=8, space=2):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), colour)
    borders.append(el)
    p_pr.append(borders)


def add_field(paragraph, instruction, placeholder="1", size=None, colour=None, bold=False):
    """Insert a Word field (PAGE, NUMPAGES, TOC ...) into a paragraph."""
    def run_with(child):
        run = paragraph.add_run()
        if size:
            run.font.size = Pt(size)
        if colour:
            run.font.color.rgb = RGBColor.from_string(colour)
        run.bold = bold
        run._r.append(child)
        return run

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run_with(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run_with(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run_with(separate)
    text_run = paragraph.add_run(placeholder)
    if size:
        text_run.font.size = Pt(size)
    if colour:
        text_run.font.color.rgb = RGBColor.from_string(colour)
    text_run.bold = bold
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


def add_rich_text(paragraph, text, size=None, colour=None, bold=False, italic=False):
    """Add text honouring **bold** and [placeholder] conventions."""
    for part in INLINE_PATTERN.split(text):
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        is_placeholder = part.startswith("[") and part.endswith("]")
        run = paragraph.add_run(part[2:-2] if is_bold else part)
        run.bold = bold or is_bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if colour:
            run.font.color.rgb = RGBColor.from_string(colour)
        if is_placeholder:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return paragraph


# --- The writer ------------------------------------------------------------------
class ProposalWriter:
    BODY_SIZE = 10
    TABLE_SIZE = 8.5
    CONTENT_WIDTH_CM = 17.0          # A4 portrait minus 2 cm margins each side
    LANDSCAPE_WIDTH_CM = 25.7

    def __init__(self):
        self.doc = Document()
        self.chapter = 0
        self.section_no = 0
        self.figure_no = 0
        self.table_no = 0
        self._configure_page(self.doc.sections[0])
        self._configure_styles()

    # -- set-up ---------------------------------------------------------------------
    def _configure_page(self, section, landscape=False):
        if landscape:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = Cm(29.7), Cm(21.0)
        else:
            section.orientation = WD_ORIENT.PORTRAIT
            section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(1.0)
        section.footer_distance = Cm(1.0)

    def _configure_styles(self):
        styles = self.doc.styles
        normal = styles["Normal"]
        _set_font(normal.font, normal.element)
        normal.font.size = Pt(self.BODY_SIZE)
        normal.font.color.rgb = RGBColor.from_string(brand.TEXT_DARK)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15

        heading_specs = {
            "Heading 1": (17, brand.MAGENTA, 0, 10, True),
            "Heading 2": (13, brand.MAGENTA, 14, 6, False),
            "Heading 3": (11, brand.ORANGE, 10, 4, False),
        }
        for name, (size, colour, before, after, page_break) in heading_specs.items():
            style = styles[name]
            _set_font(style.font, style.element)
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.italic = False
            style.font.color.rgb = RGBColor.from_string(colour)
            fmt = style.paragraph_format
            fmt.space_before = Pt(before)
            fmt.space_after = Pt(after)
            fmt.keep_with_next = True
            fmt.page_break_before = page_break

        for name in ("List Bullet", "List Bullet 2", "Caption", "Title", "Subtitle"):
            style = styles[name]
            _set_font(style.font, style.element)
        styles["List Bullet"].font.size = Pt(self.BODY_SIZE)
        styles["List Bullet"].paragraph_format.space_after = Pt(3)
        styles["List Bullet 2"].font.size = Pt(self.BODY_SIZE)
        styles["List Bullet 2"].paragraph_format.space_after = Pt(2)

        caption = styles["Caption"]
        caption.font.size = Pt(8.5)
        caption.font.bold = False
        caption.font.italic = True
        caption.font.color.rgb = RGBColor.from_string(brand.TEXT_MUTED)
        caption.paragraph_format.space_before = Pt(3)
        caption.paragraph_format.space_after = Pt(10)

        settings = self.doc.settings.element
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)

    def set_properties(self, title, subject, keywords):
        props = self.doc.core_properties
        props.title = title
        props.subject = subject
        props.author = brand.BIDDER
        props.last_modified_by = brand.BIDDER
        props.keywords = keywords
        props.category = brand.CLASSIFICATION
        props.comments = ""
        props.created = props.modified = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
        props.revision = 1

    def header_footer(self, section=None, header_text=brand.HEADER_TEXT):
        """Branded running header and 'Page X of Y' footer; blank on the cover."""
        section = section or self.doc.sections[0]
        section.different_first_page_header_footer = True

        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_rich_text(header, header_text, size=8, colour=brand.TEXT_MUTED)
        _paragraph_border(header, "bottom", brand.ORANGE, size=6)

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _paragraph_border(footer, "top", brand.MAGENTA, size=6)
        add_rich_text(footer, f"{brand.CLASSIFICATION} | Page ", size=8, colour=brand.TEXT_MUTED)
        add_field(footer, "PAGE", "1", size=8, colour=brand.TEXT_MUTED)
        add_rich_text(footer, " of ", size=8, colour=brand.TEXT_MUTED)
        add_field(footer, "NUMPAGES", "1", size=8, colour=brand.TEXT_MUTED)

    # -- headings -----------------------------------------------------------------------
    def h1(self, title, numbered=True, new_page=True):
        if numbered:
            self.chapter += 1
            self.section_no = 0
            title = f"{self.chapter}. {title}"
        heading = self.doc.add_heading(title, level=1)
        heading.paragraph_format.page_break_before = new_page
        _paragraph_border(heading, "bottom", brand.ORANGE, size=12, space=4)
        return heading

    def h2(self, title, numbered=True):
        if numbered and self.chapter:
            self.section_no += 1
            title = f"{self.chapter}.{self.section_no} {title}"
        return self.doc.add_heading(title, level=2)

    def h3(self, title):
        return self.doc.add_heading(title, level=3)

    # -- text ---------------------------------------------------------------------------
    def para(self, text, size=None, colour=None, bold=False, italic=False,
             align=None, space_after=None, keep_with_next=False):
        paragraph = self.doc.add_paragraph()
        add_rich_text(paragraph, text, size=size, colour=colour, bold=bold, italic=italic)
        if align == "center":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == "justify":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if space_after is not None:
            paragraph.paragraph_format.space_after = Pt(space_after)
        paragraph.paragraph_format.keep_with_next = keep_with_next
        return paragraph

    def paras(self, texts):
        for text in texts:
            self.para(text)

    def bullets(self, items, level=1):
        style = "List Bullet" if level == 1 else "List Bullet 2"
        for item in items:
            if isinstance(item, (list, tuple)):
                self.bullets(item, level=2)
                continue
            paragraph = self.doc.add_paragraph(style=style)
            add_rich_text(paragraph, item)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def spacer(self, points=6):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        run = paragraph.add_run()
        run.font.size = Pt(points)

    def toc(self):
        paragraph = self.doc.add_paragraph()
        add_field(paragraph, 'TOC \\o "1-2" \\h \\z \\u',
                  "Right-click here and choose Update Field to build the table of contents.")

    # -- tables -------------------------------------------------------------------------
    def table(self, headers, rows, widths, caption=None, font_size=None,
              bold_first_col=False, total_rows=0, align_right_cols=(), center_cols=(),
              header_fill=brand.MAGENTA, zebra=True, highlight_rows=(), padding=50):
        """Branded table. `rows` items may be lists of cell texts or
        ('GROUP', 'label') tuples, which render as a full-width group row.
        The last `total_rows` rows and any row index in `highlight_rows`
        are shaded and bold (used for totals and subtotals)."""
        size = font_size or self.TABLE_SIZE
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(table)
        _cell_margins(table, top=padding, bottom=padding)

        header_row = table.rows[0]
        _row_flags(header_row, header=True)
        for cell, text in zip(header_row.cells, headers):
            shade_cell(cell, header_fill)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_rich_text(paragraph, text, size=size, colour=brand.WHITE, bold=True)

        data_index = 0
        for index, row_values in enumerate(rows):
            row = table.add_row()
            _row_flags(row)
            if isinstance(row_values, tuple) and row_values and row_values[0] == "GROUP":
                merged = row.cells[0].merge(row.cells[-1])
                shade_cell(merged, brand.GROUP_ROW)
                paragraph = merged.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                add_rich_text(paragraph, row_values[1], size=size, bold=True, colour=brand.TEXT_DARK)
                data_index = 0
                continue
            is_total = index >= len(rows) - total_rows or index in highlight_rows
            for col, (cell, value) in enumerate(zip(row.cells, row_values)):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                if is_total:
                    shade_cell(cell, brand.GROUP_ROW)
                elif zebra and data_index % 2 == 1:
                    shade_cell(cell, brand.ZEBRA)
                lines = str(value).split("\n")
                for line_no, line in enumerate(lines):
                    paragraph = cell.paragraphs[0] if line_no == 0 else cell.add_paragraph()
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.05
                    if col in align_right_cols:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    elif col in center_cols:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_rich_text(paragraph, line, size=size,
                                  bold=is_total or (bold_first_col and col == 0))
            data_index += 1

        set_column_widths(table, widths)
        if caption:
            self._caption(f"Table {self._next_table()}: {caption}")
        else:
            self.spacer(4)
        return table

    def key_value_table(self, pairs, widths=(5.0, 12.0), caption=None):
        """Two-column table with a shaded label column and no header row."""
        table = self.doc.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(table)
        _cell_margins(table)
        for label, value in pairs:
            row = table.add_row()
            _row_flags(row)
            label_cell, value_cell = row.cells
            shade_cell(label_cell, brand.ZEBRA)
            for cell, text, bold in ((label_cell, label, True), (value_cell, value, False)):
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(1)
                add_rich_text(paragraph, text, size=self.TABLE_SIZE + 0.5, bold=bold)
        set_column_widths(table, widths)
        if caption:
            self._caption(f"Table {self._next_table()}: {caption}")
        else:
            self.spacer(4)
        return table

    def callout(self, title, text_lines, fill=brand.CALLOUT_FILL, accent=brand.ORANGE, width=None):
        """Highlighted box with an accent bar on the left."""
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _no_table_borders(table)
        _cell_margins(table, top=100, bottom=100, left=180, right=160)
        cell = table.rows[0].cells[0]
        shade_cell(cell, fill)
        _set_cell_borders(cell, left=(36, accent), top=None, bottom=None, right=None)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(3)
        add_rich_text(paragraph, title, size=10.5, bold=True, colour=brand.MAGENTA)
        for line in text_lines:
            item = cell.add_paragraph()
            item.paragraph_format.space_after = Pt(2)
            add_rich_text(item, line, size=self.BODY_SIZE - 0.5)
        set_column_widths(table, [width or self.CONTENT_WIDTH_CM])
        self.spacer(6)
        return table

    # -- figures ------------------------------------------------------------------------
    def figure(self, image_path, caption, width_cm=16.5):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))
        self._caption(f"Figure {self._next_figure()}: {caption}")

    def figure_grid(self, items, columns=2, max_width_cm=8.2, max_height_cm=5.6, portrait_max_height_cm=10.0):
        """Lay out (image_path, caption, aspect_ratio) items in a borderless grid,
        each image scaled to fit max_width x max_height and captioned. Portrait
        images (such as a document page) may be taller so that they stay legible."""
        table = self.doc.add_table(rows=0, cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _no_table_borders(table)
        _cell_margins(table, top=60, bottom=60, left=60, right=60)
        for start in range(0, len(items), columns):
            row = table.add_row()
            _row_flags(row)
            for cell, (path, caption, aspect) in zip(row.cells, items[start:start + columns]):
                height_limit = portrait_max_height_cm if aspect > 1 else max_height_cm
                width = min(max_width_cm, height_limit / aspect)
                picture = cell.paragraphs[0]
                picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
                picture.paragraph_format.space_after = Pt(0)
                picture.add_run().add_picture(str(path), width=Cm(width))
                label = cell.add_paragraph(style="Caption")
                label.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_rich_text(label, f"Figure {self._next_figure()}: {caption}")
        set_column_widths(table, [self.CONTENT_WIDTH_CM / columns] * columns)
        return table

    def image(self, image_path, width_cm, align="left"):
        paragraph = self.doc.add_paragraph()
        if align == "center":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))
        return paragraph

    def _caption(self, text):
        paragraph = self.doc.add_paragraph(style="Caption")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rich_text(paragraph, text)

    def _next_figure(self):
        self.figure_no += 1
        return self.figure_no

    def _next_table(self):
        self.table_no += 1
        return self.table_no

    # -- sections -----------------------------------------------------------------------
    def landscape_section(self):
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._configure_page(section, landscape=True)
        return section

    def portrait_section(self):
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._configure_page(section, landscape=False)
        return section

    def save(self, path: Path):
        self.doc.save(str(path))
