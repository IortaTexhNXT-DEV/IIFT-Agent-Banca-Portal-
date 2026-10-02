"""Shared helpers for the SalesVerse 2.0 quality and security reports.

Reuses the proposal's layout kit (docs/proposal/build) and reads evidence from
docs/technical/evidence at build time, so the reports can be rebuilt whenever
the evidence is refreshed.
"""

import hashlib
import json
import subprocess
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
TECH_DIR = ROOT / "docs" / "technical"
EVIDENCE_DIR = TECH_DIR / "evidence"
sys.path.insert(0, str(ROOT / "docs" / "proposal" / "build"))

import brand                                   # noqa: E402
import diagrams                                # noqa: E402
from docx_kit import ProposalWriter, add_rich_text   # noqa: E402
from docx.shared import Pt, Cm                 # noqa: E402

PENDING = "To be refreshed after front-end build"


def evidence_path(name: str) -> Path:
    return EVIDENCE_DIR / name


def load_json(name: str):
    """Evidence JSON, or None when the file has not been produced yet."""
    path = evidence_path(name)
    return json.loads(path.read_text()) if path.exists() else None


def load_text(name: str):
    path = evidence_path(name)
    return path.read_text().strip() if path.exists() else None


def git_commit() -> str:
    try:
        return subprocess.run(["git", "rev-parse", "--short", "HEAD"], cwd=ROOT, capture_output=True, text=True,
                              check=True).stdout.strip()
    except (subprocess.CalledProcessError, FileNotFoundError):
        return "unknown"


def format_timestamp(value) -> str:
    """ISO string or epoch milliseconds -> '2 October 2026, 14:04 UTC'."""
    if isinstance(value, (int, float)):
        moment = datetime.utcfromtimestamp(value / 1000)
    else:
        moment = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    return f"{moment.day} {moment:%B %Y, %H:%M} UTC"


def evidence_inventory():
    """(file, size, SHA-256) for every evidence file, for traceability."""
    rows = []
    for path in sorted(EVIDENCE_DIR.glob("*")):
        if path.is_file():
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            rows.append([path.name, f"{path.stat().st_size:,} bytes", digest[:16] + "…"])
    return rows


def new_report(title: str, subject: str, header_text: str) -> ProposalWriter:
    writer = ProposalWriter()
    writer.set_properties(title=title, subject=subject,
                          keywords=f"{brand.PRODUCT}; IIFT; {subject}")
    writer.header_footer(header_text=header_text)
    return writer


def cover(writer: ProposalWriter, out_dir: Path, title: str, subtitle: str, facts):
    paragraph = writer.doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(28)
    if brand.IORTA_LOGO.exists():
        paragraph.add_run().add_picture(str(brand.IORTA_LOGO), height=Cm(1.6))
    else:
        add_rich_text(paragraph, "iorta", size=26, colour=brand.MAGENTA, bold=True)
        add_rich_text(paragraph, " TechNXT", size=26, colour=brand.TEXT_DARK)
    writer.image(diagrams.cover_band(out_dir), width_cm=17.0)
    writer.spacer(18)
    writer.para(title, size=26, colour=brand.MAGENTA, bold=True, space_after=10)
    writer.para(subtitle, size=13, colour=brand.ORANGE, bold=True, space_after=36)
    writer.key_value_table(facts, widths=(5.0, 12.0))
    writer.spacer(14)
    writer.para(f"{brand.CLASSIFICATION}. Prepared by {brand.BIDDER} for {brand.CLIENT} as part of the proposal for "
                "the Agent/Banca Portal and Back-office Solution.", size=8, colour=brand.TEXT_MUTED, italic=True)


def evidence_annex(writer: ProposalWriter, title: str):
    writer.h1(title, numbered=False)
    writer.para("Results in this report are read from the evidence files below when the report is built. The "
                "fingerprint is the start of each file's SHA-256 hash, so a reader can confirm which evidence a "
                "given version of the report was built from.")
    writer.table(["Evidence file", "Size", "SHA-256 (first 16 characters)"], evidence_inventory(),
                 widths=[7.2, 3.0, 6.8], font_size=8, caption="Evidence files")


def save_and_export(writer: ProposalWriter, output: Path) -> Path:
    writer.save(output)
    from export_pdf import export      # LibreOffice UNO export; refreshes the table of contents
    return export(output)
