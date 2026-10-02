"""Build the SalesVerse 2.0 Data Dictionary (XLSX and DOCX, plus PDF of the DOCX).

Usage:
    python3 docs/technical/build/build_data_dictionary.py [--no-pdf]

Everything structural is parsed from apps/api/prisma/schema.prisma and the SQL
migrations (schema_model.py); only the business descriptions come from
dictionary_text.py. The build stops if any table or column lacks a description.
"""

import sys
import tempfile
from pathlib import Path

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parent))

import tech_kit  # noqa: E402
from tech_kit import TechnicalWriter, DocInfo, TECH_DIR, PRODUCT, brand  # noqa: E402
import tech_diagrams  # noqa: E402
import schema_model  # noqa: E402
import dictionary_text as text  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

INFO = DocInfo(
    code="DD",
    title="Data Dictionary",
    subtitle="Physical data model of the PostgreSQL database: tables, columns, keys, relationships, "
             "enumerations and database objects",
    deliverable="DEL-07 (Technical Design, database)",
    keywords="data dictionary; PostgreSQL; schema; SalesVerse 2.0; IIFT",
    purpose="This data dictionary describes every table and column of the SalesVerse 2.0 database as created by the "
            "versioned migrations in the source code. It is generated from the Prisma schema and the migration SQL, "
            "so it always matches the deployed database. A spreadsheet version with the same content is provided "
            "for filtering and searching.",
    related=["SalesVerse-2.0-Data-Dictionary.xlsx: the same content as a workbook, one sheet per table.",
             "SalesVerse-2.0-Solution-Architecture: data architecture, encryption and retention design.",
             "Source: apps/api/prisma/schema.prisma and apps/api/prisma/migrations/*/migration.sql."],
)


# --- Model preparation ---------------------------------------------------------------------------
def describe(model, column):
    key = f"{model.table}.{column.name}"
    if key in text.COLUMNS:
        return text.COLUMNS[key]
    if column.name in text.COMMON:
        return text.COMMON[column.name]
    return None


def check_descriptions(pm):
    missing = [m.table for m in pm.models if m.table not in text.TABLES]
    for model in pm.models:
        missing += [f"{model.table}.{c.name}" for c in model.columns if not describe(model, c)]
    missing += [e.name for e in pm.enums if e.name not in text.ENUMS]
    for enum in pm.enums:
        if enum.name in text.ENUMS:
            meanings = text.ENUMS[enum.name][1]
            missing += [f"{enum.name}.{v}" for v in enum.values if v not in meanings]
    missing += [name for name, _ in pm.sequences if name not in text.SEQUENCES]
    missing += [name for name, _, _ in pm.checks if name not in text.CHECKS]
    known = {f"{m.table}.{c.name}" for m in pm.models for c in m.columns}
    stale = [k for k in text.COLUMNS if k not in known]
    if missing or stale:
        raise SystemExit("Data dictionary text is out of date.\n  Missing: " + ", ".join(missing) +
                         "\n  Unknown keys: " + ", ".join(stale))


def keys_of(model, column, pm):
    parts = []
    if column.is_pk:
        parts.append("PK")
    for index in pm.indexes:
        if index.table != model.table or index.primary or column.name not in index.columns:
            continue
        label = "UQ" if index.unique else "IDX"
        if len(index.columns) > 1:
            label += f" ({index.columns.index(column.name) + 1}/{len(index.columns)})"
        parts.append(label)
    return ", ".join(dict.fromkeys(parts))


def fk_text(column):
    if not column.fk:
        return ""
    target, action = column.fk.split(" ON DELETE ")
    return f"{target}; on delete {action.lower()}"


def enum_usage(pm):
    usage = {}
    for model in pm.models:
        for column in model.columns:
            if column.is_enum:
                usage.setdefault(column.prisma_type, []).append(f"{model.table}.{column.name}")
    return usage


def index_purpose(index, pm):
    if index.primary:
        return "Primary key"
    model = pm.by_table(index.table)
    fk_cols = {c.name for c in model.columns if c.fk}
    if index.unique:
        return "Enforces uniqueness of " + " + ".join(index.columns)
    if index.columns[0] in fk_cols:
        return "Finds child rows of a parent; supports the foreign key"
    return "Supports filtering and sorting on " + ", ".join(index.columns)


def cardinality(rel):
    return "1 : 0..1" if rel.unique_child else "1 : 0..n"


def section_name(model):
    return model.section.split(" (")[0]


# --- XLSX -------------------------------------------------------------------------------------------
def build_xlsx(pm, path):
    st = tech_kit.xlsx_styles()
    wb = Workbook()
    tech_kit.set_workbook_properties(wb, f"{PRODUCT} Data Dictionary", "Physical data model")

    def sheet_header(ws, title, subtitle, headers, widths, start_row=4):
        ws["A1"] = title
        ws["A1"].font = st["title"]
        ws["A2"] = subtitle
        ws["A2"].font = st["subtitle"]
        for col, (header, width) in enumerate(zip(headers, widths), start=1):
            cell = ws.cell(row=start_row, column=col, value=header)
            cell.font, cell.fill, cell.alignment, cell.border = st["header_font"], st["header_fill"], st["wrap"], \
                st["border"]
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
        ws.sheet_view.showGridLines = False
        return start_row + 1

    def write_rows(ws, row, rows, bold_first=False, zebra=True):
        for index, values in enumerate(rows):
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.font = st["bold"] if bold_first and col == 1 else st["body"]
                cell.alignment = st["wrap"]
                cell.border = st["border"]
                if zebra and index % 2 == 1:
                    cell.fill = st["zebra_fill"]
            row += 1
        return row

    # About
    ws = wb.active
    ws.title = "About"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100
    ws["A1"] = f"{PRODUCT} Data Dictionary"
    ws["A1"].font = st["title"]
    ws["A2"] = f"{brand.SOLUTION_NAME} for {brand.CLIENT}"
    ws["A2"].font = st["subtitle"]
    about = [
        ("Document reference", f"{INFO.reference} v{INFO.version}, {tech_kit.ISSUE_DATE}"),
        ("Prepared by", brand.BIDDER),
        ("Classification", brand.CLASSIFICATION),
        ("Database", "PostgreSQL 16, schema public"),
        ("Generated from", "apps/api/prisma/schema.prisma and the SQL migrations"),
        ("Migrations included", ", ".join(pm.migrations)),
        ("Content", f"{len(pm.models)} tables, {sum(len(m.columns) for m in pm.models)} columns, "
                    f"{len(pm.enums)} enumerated types, {len(pm.relations)} foreign keys, {len(pm.indexes)} indexes "
                    f"(including primary keys), {len(pm.sequences)} sequences, {len(pm.checks)} check constraints, "
                    f"{len(pm.triggers)} trigger"),
        ("Key column legend", "PK primary key; UQ unique index; IDX non-unique index; (n/m) position in a composite "
                              "index"),
        ("Types", "TIMESTAMP(3) values are stored in UTC without time zone. DATE values are business dates in "
                  "Brunei time. DECIMAL amounts are Brunei dollars. Enumerated types are PostgreSQL ENUM types "
                  "named in CamelCase."),
        ("Defaults", "'UUID v4 generated by the application' means the id is created by the Prisma client, not by a "
                     "database default."),
        ("Protection", "See the Protection column and the Protected columns sheet for encrypted, blind-indexed and "
                       "hashed columns."),
    ]
    row = 4
    for label, value in about:
        ws.cell(row=row, column=1, value=label).font = st["bold"]
        cell = ws.cell(row=row, column=2, value=value)
        cell.font, cell.alignment = st["body"], st["wrap"]
        row += 1

    # Tables index
    ws = wb.create_sheet("Tables")
    row = sheet_header(ws, "Tables", "One row per table; the table name links to its sheet.",
                       ["Area", "Table", "Model (code)", "Columns", "Primary key", "Description"],
                       [26, 24, 22, 9, 18, 90])
    for index, model in enumerate(pm.models):
        values = [section_name(model), model.table, model.name, len(model.columns), ", ".join(model.primary_key),
                  text.TABLES[model.table]]
        row = write_rows(ws, row, [values], zebra=False)
        link = ws.cell(row=row - 1, column=2)
        link.hyperlink = f"#'{model.table}'!A1"
        link.font = st["link"]
        if index % 2 == 1:
            for col in range(1, 7):
                ws.cell(row=row - 1, column=col).fill = st["zebra_fill"]

    # One sheet per table
    headers = ["#", "Column", "Data type", "Nullable", "Default", "Keys", "Foreign key", "Protection",
               "Prisma field", "Description"]
    widths = [5, 24, 18, 9, 24, 14, 30, 26, 22, 70]
    for model in pm.models:
        ws = wb.create_sheet(model.table[:31])
        row = sheet_header(ws, model.table, text.TABLES[model.table], headers, widths, start_row=5)
        ws["A3"] = f"Area: {section_name(model)}  |  Model: {model.name}  |  Primary key: {', '.join(model.primary_key)}"
        ws["A3"].font = st["body"]
        ws["J3"] = "Back to Tables"
        ws["J3"].hyperlink = "#'Tables'!A1"
        ws["J3"].font = st["link"]
        rows = []
        for number, column in enumerate(model.columns, start=1):
            rows.append([number, column.name, column.pg_type, "Yes" if column.nullable else "No", column.default,
                         keys_of(model, column, pm), fk_text(column),
                         text.PROTECTED.get(f"{model.table}.{column.name}", ""), column.field,
                         describe(model, column)])
        write_rows(ws, row, rows)

    # Enumerations
    ws = wb.create_sheet("Enumerations")
    row = sheet_header(ws, "Enumerations", "PostgreSQL ENUM types and the meaning of each value.",
                       ["Type", "Type meaning", "Value", "Meaning", "Used by"], [22, 34, 24, 70, 50])
    usage = enum_usage(pm)
    rows = []
    for enum in pm.enums:
        meaning, values = text.ENUMS[enum.name]
        for i, value in enumerate(enum.values):
            rows.append([enum.name if i == 0 else "", meaning if i == 0 else "", value, values[value],
                         ", ".join(usage.get(enum.name, [])) if i == 0 else ""])
    write_rows(ws, row, rows, zebra=False)

    # Relationships
    ws = wb.create_sheet("Relationships")
    row = sheet_header(ws, "Relationships", "Foreign keys: the child table references the parent table.",
                       ["Constraint", "Parent table", "Parent column", "Child table", "Child column", "Cardinality",
                        "On delete", "On update", "Prisma relation field"], [40, 20, 14, 22, 20, 12, 12, 12, 22])
    rows = [[r.name, r.parent_table, ", ".join(r.parent_columns), r.child_table, ", ".join(r.child_columns),
             cardinality(r), r.on_delete, r.on_update, f"{r.child_model}.{r.field}"] for r in pm.relations]
    write_rows(ws, row, rows)

    # Indexes
    ws = wb.create_sheet("Indexes")
    row = sheet_header(ws, "Indexes", "All indexes, including primary-key indexes, as created by the migrations.",
                       ["Index", "Table", "Columns", "Unique", "Purpose"], [46, 22, 40, 9, 60])
    ordered = sorted(pm.indexes, key=lambda i: ([m.table for m in pm.models].index(i.table), not i.primary, i.name))
    rows = [[i.name, i.table, ", ".join(i.columns), "Yes" if i.unique else "No", index_purpose(i, pm)]
            for i in ordered]
    write_rows(ws, row, rows)

    # Database objects
    ws = wb.create_sheet("Database objects")
    row = sheet_header(ws, "Database objects", "Objects created by the migration SQL in addition to tables.",
                       ["Object type", "Name", "Definition", "Purpose"], [18, 34, 60, 70])
    rows = []
    for name, start in pm.sequences:
        used, example, column = text.SEQUENCES[name]
        rows.append(["Sequence", name, f"START {start}, INCREMENT 1",
                     f"{used} for {column}; formatted by NumberingService, e.g. {example}. Gaps are possible when a "
                     "transaction rolls back."])
    for name, returns, body in pm.functions:
        rows.append(["Function", name, f"RETURNS {returns}: {body}", "Raises an error for any attempt to change or "
                                                                      "delete audit rows."])
    for name, table, timing, function in pm.triggers:
        rows.append(["Trigger", name, f"{timing} ON {table} FOR EACH ROW EXECUTE FUNCTION {function}",
                     "Makes audit_log append-only at database level, whatever the client."])
    for name, table, expression in pm.checks:
        rows.append(["Check constraint", name, f"{table}: CHECK ({expression})", text.CHECKS[name]])
    write_rows(ws, row, rows)

    # Protected columns
    ws = wb.create_sheet("Protected columns")
    row = sheet_header(ws, "Protected columns", "Columns that are encrypted, blind-indexed or hashed.",
                       ["Table", "Column", "Protection", "Description"], [22, 20, 44, 90])
    rows = []
    for key, protection in text.PROTECTED.items():
        table, column = key.split(".")
        model = pm.by_table(table)
        col = next(c for c in model.columns if c.name == column)
        rows.append([table, column, protection, describe(model, col)])
    write_rows(ws, row, rows)

    summary = ["About", "Tables", "Enumerations", "Relationships", "Indexes", "Database objects",
               "Protected columns"]
    for sheet in wb.worksheets:
        if sheet.title in summary:
            sheet.sheet_properties.tabColor = brand.MAGENTA
    wb._sheets.sort(key=lambda sheet: summary.index(sheet.title) if sheet.title in summary else len(summary))
    wb.active = 0
    wb.save(path)


# --- DOCX ------------------------------------------------------------------------------------------
def build_docx(pm, path, figures):
    w = TechnicalWriter(INFO)
    w.cover(figures["cover_band"])
    w.document_control()
    w.table_of_contents()

    columns = sum(len(m.columns) for m in pm.models)
    w.h1("Introduction")
    w.h2("Scope")
    w.para(f"This dictionary covers the complete {PRODUCT} database as defined by migration "
           f"{', '.join(pm.migrations)}: {len(pm.models)} tables, {columns} columns, {len(pm.enums)} enumerated "
           f"types, {len(pm.relations)} foreign keys, {len(pm.indexes)} indexes including primary keys, "
           f"{len(pm.sequences)} sequences, {len(pm.checks)} check constraints and one trigger. Schema changes are "
           "made only through new migrations, and this document is regenerated with each release that contains one.")
    w.h2("How the dictionary is produced")
    w.para("A generator in docs/technical/build reads apps/api/prisma/schema.prisma for tables, columns, keys and "
           "relations, and the migration SQL for index names, ON DELETE rules, check constraints, sequences and the "
           "audit trigger. It compares every column type and nullability derived from the Prisma schema with the "
           "CREATE TABLE statement in the migrations and stops if they differ. Business descriptions are maintained "
           "next to the generator, and the build also stops if a table, column, enumeration value, sequence or check "
           "constraint has no description.")
    w.h2("Conventions")
    w.table(["Item", "Convention"], [
        ["Names", "Tables and columns use snake_case in PostgreSQL (from @@map and @map). The application uses "
                  "camelCase model and field names; both are listed."],
        ["Primary keys", "UUID values generated by the application, except audit_log (BIGSERIAL, to give a total "
                         "order of events), config_parameter (natural key), user_session (session id) and the two "
                         "link tables (composite keys)."],
        ["Timestamps", "TIMESTAMP(3) without time zone, written in UTC by the application. created_at defaults to "
                       "CURRENT_TIMESTAMP; updated_at is set by the application on every update."],
        ["Business dates", "DATE columns (start_date, payment_date, business_date ...) hold calendar dates in "
                           "Brunei time (UTC+8)."],
        ["Money", "DECIMAL(14,2) in Brunei dollars; rates DECIMAL(6,4); percentages DECIMAL(5,2). Amounts are never "
                  "stored as floating point."],
        ["Optimistic locking", "Tables edited by several users (agent, participant, policy) carry a version column "
                               "that is checked and incremented on update."],
        ["Keys column", "PK primary key; UQ unique index; IDX non-unique index; (n/m) is the column's position in "
                        "a composite index."],
        ["JSON", "JSONB is used for configuration and variable structures (product config, questionnaire answers, "
                 "audit before/after values, outbox payloads). Each JSON column is validated by the application "
                 "before it is written."],
    ], widths=[3.4, 13.6], bold_first_col=True, caption="Dictionary conventions")
    w.h2("Model overview")
    w.figure(figures["entities"], "Entity overview: foreign-key relationships", width_cm=16.5)
    rows = []
    current = None
    for model in pm.models:
        if section_name(model) != current:
            current = section_name(model)
            rows.append(("GROUP", current))
        rows.append([model.table, str(len(model.columns)), text.TABLES[model.table]])
    w.table(["Table", "Cols", "Description"], rows, widths=[3.8, 1.1, 12.1], font_size=8, caption="Tables by area",
            bold_first_col=True)

    # Tables in landscape for readable column descriptions.
    w.landscape_section()
    w.h1("Tables", new_page=False)
    w.para("Each table lists its columns in creation order. Foreign key shows the referenced table and column and the "
           "ON DELETE rule. Protection marks encrypted, blind-indexed or hashed values.", space_after=4)
    for model in pm.models:
        w.h2(model.table)
        w.para(f"{text.TABLES[model.table]} Model {model.name}; area {section_name(model)}; primary key "
               f"{', '.join(model.primary_key)}.", size=9, space_after=4, keep_with_next=True)
        rows = []
        for column in model.columns:
            protection = text.PROTECTED.get(f"{model.table}.{column.name}")
            description = describe(model, column)
            if protection:
                description = f"{description} **{protection}.**"
            keys = keys_of(model, column, pm)
            fk = fk_text(column)
            rows.append([column.name, column.pg_type, "Yes" if column.nullable else "No", column.default or "",
                         "\n".join(x for x in (keys, f"FK {fk}" if fk else "") if x), description])
        w.table(["Column", "Data type", "Null", "Default", "Keys / foreign key", "Description"], rows,
                widths=[3.7, 2.8, 1.0, 3.6, 4.4, 10.2], font_size=7.5, padding=30, bold_first_col=True)

    w.h1("Relationships")
    w.para("Every foreign key in the database. The parent row must exist before a child row can reference it. "
           "RESTRICT prevents deleting a parent that still has children; CASCADE deletes dependent rows with the "
           "parent; SET NULL clears the reference.")
    rows = [[r.parent_table, r.child_table, ", ".join(r.child_columns), cardinality(r), r.on_delete, r.name]
            for r in pm.relations]
    w.table(["Parent", "Child", "Child column", "Cardinality", "On delete", "Constraint"], rows,
            widths=[3.6, 4.0, 3.4, 2.4, 2.4, 9.9], font_size=8, caption="Foreign keys")
    w.para("Polymorphic references are deliberately not foreign keys: document.owner_type/owner_id, "
           "approval_request.entity_type/entity_id, aml_screening.subject_type/subject_id, "
           "outbox_message.aggregate_type/aggregate_id and audit_log.entity_type/entity_id. The application checks "
           "these references, and history tables keep the actor name so records remain readable after a user is "
           "deactivated.")

    w.h1("Indexes")
    w.para("Indexes created by the migrations, grouped by table. Composite indexes serve the most frequent list "
           "screens: for example policy (agency_id, status) for the portal policy list and "
           "outbox_message (status, next_attempt_at) for the dispatcher.")
    ordered = sorted(pm.indexes, key=lambda i: ([m.table for m in pm.models].index(i.table), not i.primary, i.name))
    rows = [[i.table, i.name, ", ".join(i.columns), "Yes" if i.unique else "No", index_purpose(i, pm)]
            for i in ordered]
    w.table(["Table", "Index", "Columns", "Unique", "Purpose"], rows, widths=[3.6, 7.6, 5.0, 1.4, 8.1],
            font_size=7.5, padding=25, caption="Indexes")

    w.h1("Enumerations")
    w.para("Enumerations are PostgreSQL ENUM types. Adding a value requires a migration; values are never removed "
           "while rows use them.")
    usage = enum_usage(pm)
    for enum in pm.enums:
        meaning, values = text.ENUMS[enum.name]
        w.h3(f"{enum.name}: {meaning[0].lower() + meaning[1:]}")
        rows = [[value, values[value]] for value in enum.values]
        w.table(["Value", "Meaning"], rows, widths=[6.0, 19.7], font_size=8, padding=25)
        w.para("Used by: " + ", ".join(usage.get(enum.name, [])), size=8, colour=brand.TEXT_MUTED, space_after=2)

    w.h1("Other Database Objects")
    w.h2("Sequences")
    w.para("Business reference numbers come from PostgreSQL sequences, which are safe under concurrency across API "
           "replicas. A rolled-back transaction can leave a gap; numbers are unique but not guaranteed contiguous.")
    rows = [[name, text.SEQUENCES[name][0], text.SEQUENCES[name][2], text.SEQUENCES[name][1]]
            for name, _ in pm.sequences]
    w.table(["Sequence", "Used for", "Column", "Format example"], rows, widths=[5.0, 6.0, 7.0, 7.7], font_size=8,
            caption="Sequences")
    w.h2("Audit trail trigger")
    for name, returns, body in pm.functions:
        w.para(f"Function **{name}** (returns {returns}) raises an exception; trigger "
               f"**{pm.triggers[0][0]}** calls it {pm.triggers[0][2]} on {pm.triggers[0][1]}, for each row. Any "
               "attempt to change or delete audit rows fails, whichever client issues it. Removing audit rows at the "
               "end of the retention period is done by a controlled archive procedure run by the database owner "
               "with the trigger disabled inside one audited maintenance transaction (delivered during "
               "implementation).")
    w.code("\n".join([
        f'CREATE FUNCTION "{pm.functions[0][0]}"() RETURNS trigger AS $$',
        "BEGIN",
        "  RAISE EXCEPTION 'audit_log is append-only';",
        "END;",
        "$$ LANGUAGE plpgsql;",
        "",
        f'CREATE TRIGGER "{pm.triggers[0][0]}"',
        f'  {pm.triggers[0][2]} ON "{pm.triggers[0][1]}"',
        f'  FOR EACH ROW EXECUTE FUNCTION "{pm.triggers[0][3]}"();',
    ]))
    w.h2("Check constraints")
    rows = [[name, table, expression, text.CHECKS[name]] for name, table, expression in pm.checks]
    w.table(["Constraint", "Table", "Expression", "Rule"], rows, widths=[6.0, 3.4, 8.6, 7.7], font_size=8,
            caption="Check constraints")
    w.h2("Session table")
    w.para("user_session has the structure expected by connect-pg-simple (sid, sess, expire). It is created by the "
           "migration, not by the library (createTableIfMissing is false), so the application database role needs no "
           "DDL rights. Expired sessions are pruned every 15 minutes.")

    w.h1("Protected Data")
    w.para("Identification numbers are never stored in clear. The application encrypts them with AES-256-GCM using "
           "FIELD_ENCRYPTION_KEY before they reach the database, with a random 96-bit IV per value; the stored form "
           "is v1:<base64 of IV, authentication tag and ciphertext>, and the version prefix allows key rotation. To "
           "find a record by identification number and to enforce uniqueness, a blind index is stored alongside: an "
           "HMAC-SHA256 of the normalised number (spaces, hyphens and slashes removed, upper case) keyed with "
           "FIELD_HASH_KEY. The unique constraints (id_type, id_number_hash) on agent and participant therefore "
           "prevent duplicates without decrypting. Screens show only the last four characters.")
    rows = []
    for key, protection in text.PROTECTED.items():
        table, column = key.split(".")
        rows.append([table, column, protection])
    w.table(["Table", "Column", "Protection"], rows, widths=[5.0, 6.0, 14.7], font_size=8,
            caption="Protected columns")
    w.para("Document files are encrypted separately with DOCUMENT_ENCRYPTION_KEY (AES-256-GCM, one IV per file); the "
           "document table holds only metadata and a generated storage key. The audit service writes [REDACTED] in "
           "place of passwords, hashes, encrypted identifiers, tokens and storage keys.")
    w.save(path)


def main():
    pdf = "--no-pdf" not in sys.argv
    pm = schema_model.load()
    check_descriptions(pm)
    TECH_DIR.mkdir(parents=True, exist_ok=True)
    xlsx = TECH_DIR / f"{INFO.file_stem}.xlsx"
    docx = TECH_DIR / f"{INFO.file_stem}.docx"
    build_xlsx(pm, xlsx)
    with tempfile.TemporaryDirectory() as tmp:
        figures = tech_diagrams.render(["entities"], Path(tmp))
        build_docx(pm, docx, figures)
    print(f"Wrote {xlsx}\nWrote {docx}")
    if pdf:
        out = tech_kit.export_pdf(docx)
        print(f"Wrote {out} ({tech_kit.pdf_pages(out)} pages)")


if __name__ == "__main__":
    main()
