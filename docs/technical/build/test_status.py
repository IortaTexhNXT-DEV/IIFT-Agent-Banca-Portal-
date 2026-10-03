"""Joins scenarios, conditions and every source of test cases into one numbered catalogue,
computes the status counts used by both the Test Strategy and the Test Cases workbook, and
verifies the references. Both documents are built from this module, and verify_pack()
reads the finished files back to confirm that neither has drifted from the other.
"""

import sys
from collections import Counter, OrderedDict
from pathlib import Path

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parent))

from reports_common import EVIDENCE_DIR, TECH_DIR, git_commit, load_json  # noqa: E402
import test_scenarios as ts  # noqa: E402
import test_automated as auto  # noqa: E402
import test_cases_manual as manual  # noqa: E402
import rfp_requirements as rfp  # noqa: E402

SCREENS_DIR = EVIDENCE_DIR / "ui-screens"
AS_OF = "2 October 2026"
STATUSES = ["Passed", "Failed", "Blocked", "Not executed"]
LEVELS = ["Unit", "API", "E2E", "UI", "Security", "SIT", "Perf", "UAT"]
LEVEL_NAMES = {
    "Unit": "Unit (API and web components)",
    "API": "API system cases executed against the running system",
    "E2E": "End-to-end API suite (business journeys)",
    "UI": "User interface (manual, exploratory)",
    "Security": "Security checks (OWASP Top 10)",
    "SIT": "System integration test with IIFT systems",
    "Perf": "Performance",
    "UAT": "User acceptance",
}

# Defects found during the test run. Severity follows the Test Strategy definitions.
DEFECTS = [
    {
        "id": "DEF-001",
        "title": "End-to-end claim test uses the UTC calendar day instead of the Brunei business day",
        "severity": "Low",
        "priority": "Medium",
        "area": "Test assets (apps/api/test/operations.e2e-spec.ts)",
        "description": "The test 'accepts claim notifications only within the period of cover' notifies a claim dated "
                       "today() in UTC. Policies seeded after 16:00 UTC start on the next Brunei day, so the claim is "
                       "refused with CLAIM_NOT_VALID and the test fails when the suite runs between 16:00 and 24:00 UTC. "
                       "The application behaves correctly (the same case passed in the API run with the Brunei date, "
                       "EX-067). The evidence run at 14:50 UTC passed; the re-run at 22:00 UTC failed.",
        "found_by": "iorta QA, re-run of the end-to-end suite on 2 October 2026",
        "status": "Open – fix in test code (use the business-day helper); no product change",
        "test": "accepts claim notifications only within the period of cover (AP-43)",
        "ex": "",
    },
    {
        "id": "DEF-002",
        "title": "Administrator password reset re-activates a disabled account",
        "severity": "Medium",
        "priority": "High",
        "area": "User administration (apps/api/src/modules/access/users.service.ts, resetPassword)",
        "description": "Resetting the password of a user whose status is DISABLED sets the status back to ACTIVE, "
                       "and the user can sign in with the temporary password (EX-134: reset 200, sign-in 200). The "
                       "audit trail records USER_PASSWORD_RESET only, not a re-activation. An administrator who "
                       "resets a password should not silently re-enable an account that another administrator "
                       "disabled; re-activation should be an explicit, separately audited status change (BO-03, "
                       "NFR-10). The unlock function behaves correctly (it refuses disabled accounts).",
        "found_by": "iorta QA, API run of 2 October 2026",
        "status": "Open – fix proposed: keep the status on reset and audit any re-activation; to be confirmed with IIFT",
        "test": "",
        "ex": "EX-134",
    },
]


def build_catalogue():
    """Returns scenarios, conditions, cases (with ids) and the API run summary."""
    executed, run_summary = auto.executed_cases()
    defect_for_test = {d["test"]: d["id"] for d in DEFECTS}
    sources = [
        auto.api_unit_cases(),
        auto.web_unit_cases(),
        auto.security_cases(),
        auto.e2e_cases(defect_for_test),
        executed,
        auto.manual_cases(SCREENS_DIR),
        auto.not_executed_cases(),
    ]
    cases = [case for group in sources for case in group]
    defect_for_ex = {d["ex"]: d["id"] for d in DEFECTS if d.get("ex")}
    for case in cases:
        if case.get("ex_id") in defect_for_ex and case["status"] == "Failed":
            case["defect"] = defect_for_ex[case["ex_id"]]
    _check_references(cases)
    area_order = {code: index for index, (code, _) in enumerate(ts.AREAS)}
    scenario_order = {s["id"]: index for index, s in enumerate(ts.SCENARIOS)}
    condition_order = {c["id"]: index for index, c in enumerate(ts.CONDITIONS)}
    level_order = {level: index for index, level in enumerate(LEVELS)}

    def sort_key(case):
        condition = ts.CONDITION_BY_ID[case["condition"]]
        scenario = ts.SCENARIO_BY_ID[condition["scenario"]]
        return (area_order[scenario["area"]], scenario_order[scenario["id"]], condition_order[condition["id"]],
                level_order.get(case["level"], 99), case["title"])

    cases.sort(key=sort_key)
    for index, case in enumerate(cases, start=1):
        condition = ts.CONDITION_BY_ID[case["condition"]]
        scenario = ts.SCENARIO_BY_ID[condition["scenario"]]
        case["id"] = f"CS-{index:04d}"
        case["scenario"] = scenario["id"]
        case["area"] = scenario["area"]
        case["area_name"] = ts.AREA_NAME[scenario["area"]]
        case["rfp"] = scenario["rfp"]
    return ts.SCENARIOS, ts.CONDITIONS, cases, run_summary


def _check_references(cases):
    problems = []
    for case in cases:
        if case["condition"] not in ts.CONDITION_BY_ID:
            problems.append(f"unknown condition {case['condition']} in case '{case['title']}'")
        if case["status"] not in STATUSES:
            problems.append(f"unknown status {case['status']} in case '{case['title']}'")
        if case["level"] not in LEVELS:
            problems.append(f"unknown level {case['level']} in case '{case['title']}'")
        if not case["expected"].strip():
            problems.append(f"case without expected result: '{case['title']}'")
    for condition in ts.CONDITIONS:
        if condition["scenario"] not in ts.SCENARIO_BY_ID:
            problems.append(f"condition {condition['id']} refers to unknown scenario {condition['scenario']}")
    used_conditions = {case["condition"] for case in cases}
    for condition in ts.CONDITIONS:
        if condition["id"] not in used_conditions:
            problems.append(f"condition {condition['id']} has no test case")
    used_scenarios = {c["scenario"] for c in ts.CONDITIONS}
    for scenario in ts.SCENARIOS:
        if scenario["id"] not in used_scenarios:
            problems.append(f"scenario {scenario['id']} has no condition")
        for rid in [r.strip() for r in scenario["rfp"].split(",")]:
            if rid not in rfp.REQUIREMENT_BY_ID:
                problems.append(f"scenario {scenario['id']} cites unknown requirement {rid}")
    if len({s["id"] for s in ts.SCENARIOS}) != len(ts.SCENARIOS) or len({c["id"] for c in ts.CONDITIONS}) != len(ts.CONDITIONS):
        problems.append("duplicate scenario or condition identifier")
    for defect in DEFECTS:
        if defect["test"] and not any(defect["test"] == entry["test"] for entry in manual.E2E):
            problems.append(f"defect {defect['id']} refers to an unknown test")
        if defect.get("ex") and defect["ex"] not in manual.EXECUTED:
            problems.append(f"defect {defect['id']} refers to an unknown executed case")
    for case in cases:
        if case["status"] == "Failed" and not case["defect"]:
            problems.append(f"failed case without a defect reference: '{case['title']}'")
    if problems:
        raise SystemExit("Test catalogue inconsistent:\n  " + "\n  ".join(problems))


def summary(cases):
    """Counts shared by the strategy and the workbook."""
    by_status = OrderedDict((status, sum(1 for c in cases if c["status"] == status)) for status in STATUSES)
    by_level = OrderedDict()
    for level in LEVELS:
        subset = [c for c in cases if c["level"] == level]
        by_level[level] = OrderedDict([("total", len(subset))] + [(s, sum(1 for c in subset if c["status"] == s)) for s in STATUSES])
    by_area = OrderedDict()
    for code, name in ts.AREAS:
        subset = [c for c in cases if c["area"] == code]
        by_area[code] = OrderedDict([("name", name), ("total", len(subset))] + [(s, sum(1 for c in subset if c["status"] == s)) for s in STATUSES])
    by_automation = Counter("Automated" if c["automation"].startswith("Automated") else
                            "Executed – API run" if c["automation"].startswith("Executed") else
                            "Manual" if c["automation"].startswith("Manual") else "Not automated" for c in cases)
    return {
        "scenarios": len(ts.SCENARIOS),
        "conditions": len(ts.CONDITIONS),
        "cases": len(cases),
        "by_status": by_status,
        "by_level": by_level,
        "by_area": by_area,
        "by_automation": OrderedDict((k, by_automation.get(k, 0)) for k in ("Automated", "Executed – API run", "Manual", "Not automated")),
        "defects": len(DEFECTS),
        "defects_open": sum(1 for d in DEFECTS if d["status"].startswith("Open")),
        "as_of": AS_OF,
        "commit": git_commit(),
    }


# Requirements met by documents or design rather than by a test case.
REVIEW_ONLY = {
    "NFR-20": "Verified by review of the Solution Architecture (modular monolith); not a test case",
    "NFR-21": "Verified by the document pack and its rebuild from source; not a test case",
    "DEL-12": "This Test Strategy",
    "DEL-13": "This workbook (Scenarios, Conditions, Cases)",
    "DEL-14": "Execution log and evidence files of this workbook; the SIT results report follows SIT",
}


def traceability(cases):
    """One row per RFP requirement: scenarios, conditions, case ids and coverage."""
    rows = []
    cases_by_scenario = {}
    for case in cases:
        cases_by_scenario.setdefault(case["scenario"], []).append(case)
    for rid, area, title in rfp.REQUIREMENTS:
        scenarios = [s for s in ts.SCENARIOS if rid in [r.strip() for r in s["rfp"].split(",")]]
        linked = [c for s in scenarios for c in cases_by_scenario.get(s["id"], [])]
        executed = [c for c in linked if c["status"] in ("Passed", "Failed")]
        if not scenarios and rid in REVIEW_ONLY:
            coverage = REVIEW_ONLY[rid]
        elif not scenarios:
            coverage = "No test case yet"
        elif not linked:
            coverage = "Scenario defined, no case yet"
        elif not executed:
            coverage = "Cases defined, not executed"
        elif any(c["status"] == "Failed" for c in executed):
            coverage = "Executed – failure open"
        elif len(executed) == len(linked):
            coverage = "Executed – passed"
        else:
            coverage = "Partly executed – passed"
        rows.append({
            "id": rid, "area": area, "title": title,
            "scenarios": ", ".join(s["id"] for s in scenarios),
            "cases": len(linked), "executed": len(executed),
            "passed": sum(1 for c in executed if c["status"] == "Passed"),
            "failed": sum(1 for c in executed if c["status"] == "Failed"),
            "case_ids": _compact_ids([c["id"] for c in linked]),
            "coverage": coverage,
        })
    return rows


def _compact_ids(ids):
    """CS-0001, CS-0002, CS-0003, CS-0007 -> CS-0001–CS-0003, CS-0007."""
    numbers = sorted(int(i.split("-")[1]) for i in ids)
    ranges, start, previous = [], None, None
    for n in numbers:
        if start is None:
            start = previous = n
        elif n == previous + 1:
            previous = n
        else:
            ranges.append((start, previous))
            start = previous = n
    if start is not None:
        ranges.append((start, previous))
    return ", ".join(f"CS-{a:04d}" if a == b else f"CS-{a:04d}–CS-{b:04d}" for a, b in ranges)


# --- Cross-document verification ------------------------------------------------------------
XLSX = TECH_DIR / "SalesVerse-2.0-Test-Cases.xlsx"
DOCX = TECH_DIR / "SalesVerse-2.0-Test-Strategy.docx"


def verify_pack(expected=None):
    """Fails when the workbook's recalculated Summary or the strategy's status table differ from the catalogue."""
    import sys as _sys
    _sys.path.insert(0, str(TECH_DIR.parent / "proposal" / "build"))
    from xlsx_kit import recalculated_values, recalc_errors
    from docx import Document

    if expected is None:
        _, _, cases, _ = build_catalogue()
        expected = summary(cases)
    failures = []
    values = recalculated_values(XLSX)
    if values is None:
        failures.append("LibreOffice not available to recalculate the workbook")
    else:
        ws = values["Summary"]
        found = {}
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            label = row[0].value
            if isinstance(label, str) and label.startswith("Status:"):
                found[label.split(":", 1)[1].strip()] = row[1].value
            if label == "Test cases":
                found["cases"] = row[1].value
            if label == "Test scenarios":
                found["scenarios"] = row[1].value
            if label == "Test conditions":
                found["conditions"] = row[1].value
        for status, count in expected["by_status"].items():
            if found.get(status) != count:
                failures.append(f"workbook Summary '{status}' = {found.get(status)}, catalogue = {count}")
        for key in ("cases", "scenarios", "conditions"):
            if found.get(key) != expected[key]:
                failures.append(f"workbook Summary {key} = {found.get(key)}, catalogue = {expected[key]}")
        errors = recalc_errors(XLSX)
        if errors:
            failures.append(f"workbook error cells: {errors[:10]}")
    document = Document(str(DOCX))
    table = next((t for t in document.tables if t.rows[0].cells[0].text.strip() == "Status" and "Cases" in t.rows[0].cells[1].text), None)
    if table is None:
        failures.append("strategy document has no status table")
    else:
        doc_counts = {row.cells[0].text.strip(): row.cells[1].text.strip() for row in table.rows[1:]}
        for status, count in expected["by_status"].items():
            if doc_counts.get(status) != str(count):
                failures.append(f"strategy '{status}' = {doc_counts.get(status)}, catalogue = {count}")
        if doc_counts.get("Total") != str(expected["cases"]):
            failures.append(f"strategy total = {doc_counts.get('Total')}, catalogue = {expected['cases']}")
    if failures:
        raise AssertionError("Test documents have drifted:\n  " + "\n  ".join(failures))
    return expected


if __name__ == "__main__":
    scenarios, conditions, cases, run = build_catalogue()
    s = summary(cases)
    print(f"{s['scenarios']} scenarios, {s['conditions']} conditions, {s['cases']} cases")
    print(dict(s["by_status"]))
    for level, counts in s["by_level"].items():
        print(f"  {level:9} {dict(counts)}")
    print(dict(s["by_automation"]))
    rows = traceability(cases)
    print(Counter(r["coverage"] for r in rows))
